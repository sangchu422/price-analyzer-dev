from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.analysis.models import QuoteAnalysisLineResult, QuoteAnalysisRun
from app.analysis import family_analysis
from app.analysis.service import AnalysisLine, AnalysisSource, DocumentAnalysis
from app.analysis.target_price import AnalysisRunResult, TargetLineResult
from app.db.base import Base
from app.db.sqlite import configure_sqlite
from app.documents.models import SourceDocument, SourceVariant
from app.procurement import equipment
from app.procurement.categories import classify_category
from app.procurement.equipment import EquipmentDefinition
from app.procurement.models import QuoteAnalysisEquipmentGroup
from app.quotes.models import RawQuoteItem


def _analysis_line(
    raw_id: int,
    *,
    sheet: str,
    quote_amount: str,
    path: str = "quote.xlsx",
) -> AnalysisLine:
    return AnalysisLine(
        raw_item_id=raw_id,
        item_name=f"ITEM {raw_id}",
        spec=f"SPEC-{raw_id}",
        spec_source_status="PRESENT",
        unit="EA",
        quantity=Decimal("1"),
        quote_unit_price=Decimal(quote_amount),
        quote_amount=Decimal(quote_amount),
        match_status="MATCHED",
        assessment="HIGH",
        reference_price=Decimal("100"),
        minimum_price=Decimal("100"),
        average_price=Decimal("100"),
        maximum_price=Decimal("100"),
        variance_amount=Decimal("1"),
        variance_percent=Decimal("1"),
        clean_decision_id=None,
        membership_decision_id=None,
        standard_item_id=None,
        standard_item_version_id=None,
        canonical_name=None,
        canonical_spec=None,
        canonical_unit=None,
        standard_price_version_id=None,
        standard_price_item_version_id=None,
        standard_observation_count=1,
        evidence_quality="SINGLE_OBSERVATION",
        market_price_lookup_required=False,
        market_price_lookup_status="NOT_REQUIRED",
        candidates=(),
        source=AnalysisSource(
            document_id=1,
            logical_name="quote.xlsx",
            variant_id=1,
            path=path,
            sha256="a" * 64,
            sheet=sheet,
            page=None,
            row=raw_id,
            cells=f"A{raw_id}:I{raw_id}",
            parser_name="xlsx",
            parser_version="reader-v2",
        ),
    )


def _target_line(raw_id: int, target_amount: str) -> TargetLineResult:
    return TargetLineResult(
        raw_item_id=raw_id,
        status="AVAILABLE",
        target_unit_price=Decimal(target_amount),
        target_amount=Decimal(target_amount),
        variance_amount=Decimal("0"),
        variance_percent=Decimal("0"),
        used_observation_count=1,
        excluded_observation_count=0,
        reason="test",
        evidence=(),
        unit_variance_amount=Decimal("0"),
    )


def test_category_rules_keep_search_grouping_separate_from_price_matching() -> None:
    assert classify_category("SERVO MOTOR", "1KW").code == "DRIVE_MOTION"
    assert classify_category("AIR CYLINDER", "25mm").code == "PNEUMATIC_HYDRAULIC"
    assert classify_category("NETWORK SWITCH", "8 PORT").code == "IT_NETWORK"
    assert classify_category("PALLET STOPPER", None).code == "MATERIAL_HANDLING"
    fallback = classify_category("SPECIAL ASSEMBLY", "ZX-991")
    assert fallback.code == "GENERAL_COMPONENT"
    assert fallback.confidence == Decimal("25")


def test_family_analysis_uses_family_minimum_and_sheet_totals(monkeypatch) -> None:
    family = {
        "code": "OTHER_GENERAL_COMPONENT",
        "name": "공통 설비·부품류",
        "item_count": 4,
        "observation_count": 9,
        "supplier_count": 3,
        "price": {
            "minimum": "80",
            "median": "90",
            "average": "92",
            "maximum": "120",
        },
        "members": [
            {
                "standard_item_id": 12,
                "name": "ITEM",
                "spec": "SPEC",
                "unit": "EA",
                "observation_count": 3,
                "price": {"minimum": "75", "median": "80", "average": "82", "maximum": "90"},
            }
        ],
    }
    monkeypatch.setattr(family_analysis, "item_family_projection", lambda _session: [family])
    line = _analysis_line(1, sheet="설비1", quote_amount="100")
    result = AnalysisRunResult(
        run_id=1,
        analysis=DocumentAnalysis(
            document_id=1,
            logical_name="quote.xlsx",
            lines=(line,),
            next_cursor=None,
            limit=100,
        ),
        inflation_sync_run_id=None,
        inflation_series_kind=None,
        target_period=None,
        target_index_value=None,
        inflation_source_url="",
        inflation_source_last_changed=None,
        quote_total_amount=Decimal("100"),
        target_total_amount=None,
        target_available_count=0,
        target_unavailable_count=1,
        target_lines=(_target_line(1, "100"),),
    )

    payload = family_analysis.family_analysis_payload(
        None,
        result,
        review_percent=Decimal("10"),
        high_percent=Decimal("20"),
    )

    assert payload["matched_count"] == 1
    assert payload["target_lines"][0]["target_unit_price"] == Decimal("80")
    assert payload["target_lines"][0]["calculation_basis"] == {
        "kind": "ITEM_FAMILY",
        "family_code": "OTHER_GENERAL_COMPONENT",
        "family_name": "공통 설비·부품류",
        "unit": "EA",
        "price_band_low": Decimal("70.0"),
        "price_band_high": Decimal("130.0"),
        "candidate_item_count": 1,
        "selection_rule": "비교군 상세품목별 중앙값 중 현재 단가 이하 최저값",
    }
    assert payload["target_lines"][0]["comparison_evidence"][0]["selected"] is True
    assert payload["equipment_groups"][0]["name"] == "설비1"
    assert payload["equipment_groups"][0]["quote_amount"] == Decimal("100")
    assert payload["equipment_groups"][0]["target_amount"] == Decimal("80")
    assert payload["equipment_groups"][0]["negotiation_amount"] == Decimal("20")


def test_equipment_definitions_resolve_digest_relative_submission_path(
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submissions"
    workbook_path = submission_root / "digest" / "quote.xlsx"
    workbook_path.parent.mkdir(parents=True)
    workbook = Workbook()
    cover = workbook.active
    cover.title = "갑지"
    cover["A1"] = "단위: 천원"
    cover["C5"] = "백래쉬시험기"
    cover["H5"] = 262291.7703
    detail = workbook.create_sheet("단위장비1")
    detail["C1"] = "백래쉬시험기"
    workbook.save(workbook_path)
    monkeypatch.setattr(equipment.settings, "submission_folder", submission_root)
    line = _analysis_line(
        1,
        sheet="단위장비1",
        quote_amount="100",
        path="digest/quote.xlsx",
    )
    result = AnalysisRunResult(
        run_id=1,
        analysis=DocumentAnalysis(
            document_id=1,
            logical_name="quote.xlsx",
            lines=(line,),
            next_cursor=None,
            limit=100,
        ),
        inflation_sync_run_id=None,
        inflation_series_kind=None,
        target_period=None,
        target_index_value=None,
        inflation_source_url="",
        inflation_source_last_changed=None,
        quote_total_amount=Decimal("100"),
        target_total_amount=Decimal("90"),
        target_available_count=1,
        target_unavailable_count=0,
        target_lines=(_target_line(1, "90"),),
    )

    definitions = equipment._equipment_definitions(result)

    assert definitions == (
        EquipmentDefinition(
            "equipment-1",
            "백래쉬시험기",
            "단위장비1",
            Decimal("262291770"),
            "COVER_SHEET",
        ),
    )


def test_equipment_projection_sums_detail_sheet_and_ignores_empty_cover_group(
    tmp_path: Path,
    monkeypatch,
) -> None:
    engine = configure_sqlite(
        create_engine(f"sqlite:///{(tmp_path / 'equipment.sqlite3').as_posix()}")
    )
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        document = SourceDocument(logical_name="quote.xlsx")
        variant = SourceVariant(
            document=document,
            path="quote.xlsx",
            sha256="a" * 64,
            extension=".xlsx",
            security_state="UNLOCKED",
            selected_for_parsing_at_ingest=True,
        )
        raw_rows = [
            RawQuoteItem(
                source_variant=variant,
                source_sheet="단위장비1",
                source_row=index,
                source_cells=f"A{index}:I{index}",
                item_name_raw=f"ITEM {index}",
                parser_name="xlsx",
                parser_version="reader-v2",
            )
            for index in (1, 2)
        ]
        session.add_all([document, *raw_rows])
        session.flush()
        run = QuoteAnalysisRun(
            document_id=document.id,
            created_by="buyer",
            review_percent=Decimal("10"),
            high_percent=Decimal("20"),
            inflation_sync_run_id=None,
            target_period=None,
            target_index_value=None,
            total_line_count=2,
            target_available_count=2,
            target_unavailable_count=0,
            quote_total_amount=Decimal("1300"),
            target_total_amount=Decimal("1000"),
        )
        session.add(run)
        session.flush()
        for raw in raw_rows:
            session.add(
                QuoteAnalysisLineResult(
                    analysis_run_id=run.id,
                    raw_item_id=raw.id,
                    standard_price_version_id=None,
                    match_status="MATCHED",
                    assessment="HIGH",
                    reference_price=Decimal("100"),
                    variance_amount=Decimal("1"),
                    variance_percent=Decimal("1"),
                    target_status="AVAILABLE",
                    target_unit_price=Decimal("1"),
                    target_amount=Decimal("1"),
                    target_variance_amount=Decimal("1"),
                    target_variance_percent=Decimal("1"),
                    target_unit_variance_amount=Decimal("1"),
                    target_used_observation_count=1,
                    target_excluded_observation_count=0,
                    target_reason="test",
                )
            )
        session.flush()
        lines = (
            _analysis_line(raw_rows[0].id, sheet="단위장비1", quote_amount="600"),
            _analysis_line(raw_rows[1].id, sheet="단위장비1", quote_amount="700"),
        )
        targets = (
            _target_line(raw_rows[0].id, "400"),
            _target_line(raw_rows[1].id, "600"),
        )
        result = AnalysisRunResult(
            run_id=run.id,
            analysis=DocumentAnalysis(
                document_id=document.id,
                logical_name=document.logical_name,
                lines=lines,
                next_cursor=None,
                limit=100,
            ),
            inflation_sync_run_id=None,
            inflation_series_kind=None,
            target_period=None,
            target_index_value=None,
            inflation_source_url="",
            inflation_source_last_changed=None,
            quote_total_amount=Decimal("1300"),
            target_total_amount=Decimal("1000"),
            target_available_count=2,
            target_unavailable_count=0,
            target_lines=targets,
        )
        monkeypatch.setattr(
            equipment,
            "_equipment_definitions",
            lambda _result: (
                EquipmentDefinition(
                    "equipment-1",
                    "백래쉬시험기",
                    "단위장비1",
                    Decimal("1000"),
                    "COVER_SHEET",
                ),
                EquipmentDefinition(
                    "equipment-2",
                    "비틀림시험기",
                    "단위장비2",
                    Decimal("600"),
                    "COVER_SHEET",
                ),
            ),
        )

        groups = equipment.create_equipment_projection(session, result)
        assert [(group.equipment_name, group.quote_amount) for group in groups] == [
            ("백래쉬시험기", Decimal("1300")),
        ]
        assert groups[0].negotiation_amount == Decimal("300")
        assert groups[0].target_amount == Decimal("1000")
        assert groups[0].unallocated_amount == Decimal("0")
        assert len(
            session.scalars(select(QuoteAnalysisEquipmentGroup)).all()
        ) == 1
        fallback = equipment.equipment_group_payloads_from_result(result)
        assert [(row["name"], row["quote_amount"]) for row in fallback] == [
            ("백래쉬시험기", Decimal("1300")),
        ]
        assert fallback[0]["target_amount"] == Decimal("1000")
        assert fallback[0]["negotiation_amount"] == Decimal("300")

        monkeypatch.setattr(equipment, "_equipment_tables_available", lambda _session: False)
        assert equipment.create_equipment_projection(session, result) == ()
        assert equipment.equipment_group_payloads(session, run.id) == []
    engine.dispose()
