from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.catalog.models import DocumentMetadataVersion
from app.db.base import Base
from app.db.sqlite import configure_sqlite
from app.documents.models import SourceDocument, SourceVariant
from app.metadata_audit.team_standard_dates import (
    TeamStandardDateError,
    backfill_team_standard_dates,
)
from app.standard_database.models import (
    QuoteDocumentPurpose,
    QuoteDocumentRole,
)


def _write_team_sources(
    tmp_path: Path,
    *,
    source_file: str,
    quote_date: str,
    json_quote_date: str | None = None,
) -> tuple[Path, Path]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    workbook_path = tmp_path / "표준단가DB.xlsx"
    source_json_path = tmp_path / "표준단가DB_전체원본.json"
    workbook = Workbook()
    standard = workbook.active
    standard.title = "표준단가DB"
    standard.append(["title"])
    standard.append(["description"])
    standard.append([
        "표준품목ID", "품목분류", "품명", "규격", "단위",
        "단가_최저(원)", "단가_평균(원)", "단가_최고(원)",
        "데이터건수", "주요메이커", "최근견적일", "출처공급사",
        "데이터입력일자",
    ])
    standard.append([
        "STD-0001", "설비", "MOTOR", "1KW", "EA",
        100, 100, 100, 1, "MAKER", quote_date, "SUPPLIER", "2026-08-11",
    ])
    raw = workbook.create_sheet("원본견적데이터")
    raw.append([
        "ID", "공급사", "견적번호", "견적일", "공사명", "단위장비명",
        "품목분류", "품명", "규격", "단위", "수량", "단가(원)",
        "금액(원)", "메이커", "데이터입력일자",
    ])
    raw.append([
        1, "SUPPLIER", "Q-001", quote_date, "PROJECT", "DEVICE",
        "설비", "MOTOR", "1KW", "EA", 1, 100, 100, "MAKER", "2026-08-11",
    ])
    workbook.save(workbook_path)
    source_json_path.write_text(
        json.dumps([
            {
                "단위장비명": "DEVICE",
                "품명": "MOTOR",
                "규격": "1KW",
                "단위": "EA",
                "수량": 1,
                "단가_원": 100,
                "금액_원": 100,
                "메이커": "MAKER",
                "ID": 1,
                "공급사": "SUPPLIER",
                "견적번호": "Q-001",
                "견적일": json_quote_date or quote_date,
                "공사명": "PROJECT",
                "품목분류": "설비",
                "출처파일": source_file,
                "데이터입력일자": "2026-08-11",
            }
        ], ensure_ascii=False),
        encoding="utf-8",
    )
    return workbook_path, source_json_path


def _add_variant(
    session: Session,
    path: str,
) -> SourceVariant:
    document = SourceDocument(logical_name=path)
    session.add(document)
    session.flush()
    variant = SourceVariant(
        document_id=document.id,
        path=path,
        sha256="a" * 64,
        extension=Path(path).suffix,
        security_state="OPEN",
        selected_for_parsing_at_ingest=True,
    )
    session.add_all([
        variant,
        QuoteDocumentRole(
            document_id=document.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            supersedes_role_id=None,
            decided_by="test",
            reason_detail="historical source",
        ),
    ])
    session.flush()
    return variant


def _session() -> tuple[object, Session]:
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    return engine, Session(engine)


def test_backfill_adds_inferred_date_and_is_idempotent(tmp_path: Path) -> None:
    source = (
        "바츠 추출 견적서/tmp_20260803/"
        "20240115120000000source.pdf"
    )
    workbook, source_json = _write_team_sources(
        tmp_path,
        source_file=source,
        quote_date="2024-01-15",
    )
    engine, session = _session()
    try:
        variant = _add_variant(session, f"3차 학습/{source}")

        first = backfill_team_standard_dates(
            session,
            workbook_path=workbook,
            source_json_path=source_json,
            source_revision="team/main@test",
            apply=True,
        )
        second = backfill_team_standard_dates(
            session,
            workbook_path=workbook,
            source_json_path=source_json,
            source_revision="team/main@test",
            apply=True,
        )
        session.commit()

        assert first.applied_files == 1
        assert second.applied_files == 0
        assert second.already_same_files == 1
        assert session.scalar(
            select(func.count(DocumentMetadataVersion.id))
        ) == 1
        metadata = session.scalar(select(DocumentMetadataVersion))
        assert metadata is not None
        assert metadata.source_document_id == variant.document_id
        assert metadata.quote_date == date(2024, 1, 15)
        evidence = json.loads(metadata.evidence_json)["quote_date"]
        assert evidence["quality"] == "FILE_DATE_INFERRED"
        assert evidence["use_for_index"] == "YEAR_ONLY"
    finally:
        session.close()
        engine.dispose()


def test_backfill_rejects_placeholder_and_preserves_confirmed_date(
    tmp_path: Path,
) -> None:
    workbook, source_json = _write_team_sources(
        tmp_path / "placeholder",
        source_file="placeholder.xlsx",
        quote_date="2025-01-01",
    )
    conflict_workbook, conflict_json = _write_team_sources(
        tmp_path / "conflict",
        source_file="conflict.xlsx",
        quote_date="2025-05-20",
    )
    engine, session = _session()
    try:
        _add_variant(session, "1차 학습/placeholder.xlsx")
        conflict = _add_variant(session, "1차 학습/conflict.xlsx")
        session.add(
            DocumentMetadataVersion(
                source_document_id=conflict.document_id,
                version_number=1,
                supplier_name=None,
                quote_date=date(2025, 5, 10),
                project_name=None,
                decided_by="metadata-audit-v2",
                reason_detail="source label",
                evidence_json="{}",
            )
        )
        session.flush()

        placeholder_report = backfill_team_standard_dates(
            session,
            workbook_path=workbook,
            source_json_path=source_json,
            source_revision="team/main@test",
            apply=True,
        )
        conflict_report = backfill_team_standard_dates(
            session,
            workbook_path=conflict_workbook,
            source_json_path=conflict_json,
            source_revision="team/main@test",
            apply=True,
        )

        assert placeholder_report.placeholder_files == 1
        assert placeholder_report.applied_files == 0
        assert conflict_report.conflict_files == 1
        assert conflict_report.applied_files == 0
        current = session.scalar(
            select(DocumentMetadataVersion).where(
                DocumentMetadataVersion.source_document_id == conflict.document_id
            )
        )
        assert current is not None
        assert current.quote_date == date(2025, 5, 10)
    finally:
        session.close()
        engine.dispose()


def test_backfill_requires_workbook_json_row_equality(tmp_path: Path) -> None:
    workbook, source_json = _write_team_sources(
        tmp_path,
        source_file="source.xlsx",
        quote_date="2025-05-20",
        json_quote_date="2025-05-21",
    )
    engine, session = _session()
    try:
        with pytest.raises(TeamStandardDateError, match="differ at raw ID"):
            backfill_team_standard_dates(
                session,
                workbook_path=workbook,
                source_json_path=source_json,
                source_revision="team/main@test",
                apply=False,
            )
    finally:
        session.close()
        engine.dispose()
