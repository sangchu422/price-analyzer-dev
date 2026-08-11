from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.catalog.models import (
    DocumentMetadataCandidate,
    DocumentMetadataScan,
    DocumentMetadataVersion,
)
from app.db.base import Base
from app.db.sqlite import configure_sqlite
from app.ingestion.service import ingest_path
from app.metadata_audit.service import audit_quote_metadata
from app.metadata_audit.service import (
    ExtractedCandidate,
    _date_text_candidates,
    _select_candidates,
)
from app.standard_database.service import assign_initial_historical_roles


def _write_quote(path: Path, supplier: str, quote_date: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "견적"
    sheet.append(["공급사", supplier])
    sheet.append(["견적일", quote_date])
    sheet.append(["공사명", "자동화 라인"])
    sheet.append([])
    sheet.append(["품명", "규격", "단위", "수량", "단가", "금액"])
    sheet.append(["SERVO MOTOR", "1KW", "EA", 1, 1000000, 1000000])
    workbook.save(path)


def test_audit_uses_labeled_source_values_not_collection_folder(tmp_path: Path) -> None:
    quote_root = tmp_path / "견적서"
    quote = quote_root / "3차 학습" / "AONE 추출 견적서" / "actual.xlsx"
    _write_quote(quote, "실제공급사", "2024-05-20")
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        ingest_path(session, quote, root=quote_root)
        report_path = tmp_path / "audit.csv"

        first = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=report_path,
        )
        second = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=report_path,
        )
        session.commit()

        assert first.total_files == second.total_files == 1
        assert first.auto_confirmed_files == 1
        assert session.scalar(select(func.count(DocumentMetadataScan.id))) == 1
        assert session.scalar(select(func.count(DocumentMetadataCandidate.id))) == 3
        metadata = session.scalar(select(DocumentMetadataVersion))
        assert metadata is not None
        assert metadata.supplier_name == "실제공급사"
        assert metadata.quote_date.isoformat() == "2024-05-20"
        assert metadata.project_name == "자동화 라인"
        assert "AONE" not in metadata.supplier_name
        with report_path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        assert rows[0]["acquisition_channel"] == "AONE"
        assert rows[0]["supplier_name"] == "실제공급사"


def test_audit_reads_unlabeled_legal_company_name_from_quote_cover(
    tmp_path: Path,
) -> None:
    quote_root = tmp_path / "견적서"
    quote = quote_root / "3차 학습" / "cover-company.xlsx"
    quote.parent.mkdir(parents=True)
    workbook = Workbook()
    cover = workbook.active
    cover.title = "갑지"
    cover["B2"] = "견 적 서"
    cover["G4"] = "한로기술(주)"
    cover["B5"] = "견적일자"
    cover["D5"] = "2025-11-15"
    detail = workbook.create_sheet("단위장비4")
    detail.append(["품명", "규격", "단위", "수량", "단가", "금액"])
    detail.append(["2분력계", "5000NM/10000N", "SET", 2, 47000000, 94000000])
    workbook.save(quote)
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        ingest_path(session, quote, root=quote_root)
        report = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=tmp_path / "cover-audit.csv",
        )
        session.commit()

        metadata = session.scalar(select(DocumentMetadataVersion))
        assert report.auto_confirmed_files == 1
        assert metadata is not None
        assert metadata.supplier_name == "한로기술(주)"
        assert metadata.quote_date.isoformat() == "2025-11-15"
        assert '"cells":"G4"' in metadata.evidence_json


def test_audit_rejects_collection_labels_and_placeholder_dates(tmp_path: Path) -> None:
    quote_root = tmp_path / "견적서"
    quote = quote_root / "3차 학습" / "바츠 추출 견적서" / "placeholder.xlsx"
    _write_quote(quote, "바츠추출", "2025-01-01")
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        ingest_path(session, quote, root=quote_root)

        report = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=tmp_path / "audit.csv",
        )
        session.commit()

        assert report.review_required_files == 1
        assert session.scalar(select(func.count(DocumentMetadataVersion.id))) == 1
        metadata = session.scalar(select(DocumentMetadataVersion))
        assert metadata is not None
        assert metadata.supplier_name is None
        assert metadata.quote_date is None
        assert metadata.project_name == "자동화 라인"


def test_audit_records_unregistered_manual_review_files(tmp_path: Path) -> None:
    quote_root = tmp_path / "견적서"
    audit_root = quote_root / "3차 학습"
    quote = audit_root / "source.xlsx"
    image = audit_root / "scan.jpg"
    _write_quote(quote, "실제공급사", "2024-05-20")
    image.write_bytes(b"manual image review")
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        ingest_path(session, quote, root=quote_root)

        report = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=tmp_path / "audit.csv",
        )
        session.commit()

        assert report.total_files == 2
        assert report.unsupported_files == 1
        scans = list(session.scalars(select(DocumentMetadataScan)))
        assert len(scans) == 2
        by_name = {Path(scan.source_path).name: scan for scan in scans}
        assert by_name["source.xlsx"].source_variant_id is not None
        assert by_name["scan.jpg"].source_variant_id is None
        assert by_name["scan.jpg"].review_status == "REVIEW_REQUIRED"


def test_audit_allows_corpus_without_third_training_folder(
    tmp_path: Path,
) -> None:
    quote_root = tmp_path / "견적서"
    quote_root.mkdir()
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    report_path = tmp_path / "audit.csv"

    with Session(engine) as session:
        report = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=report_path,
        )

    assert report.total_files == 0
    assert report.review_required_files == 0
    with report_path.open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.DictReader(stream)) == []


def test_all_historical_audit_includes_sources_outside_third_training(
    tmp_path: Path,
) -> None:
    quote_root = tmp_path / "견적서"
    quote = quote_root / "1차 학습" / "actual.xlsx"
    _write_quote(quote, "실제공급사", "2024-05-20")
    engine = configure_sqlite(create_engine("sqlite:///:memory:"))
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        ingest_path(session, quote, root=quote_root)
        assign_initial_historical_roles(session, actor="test")

        report = audit_quote_metadata(
            session,
            quote_root=quote_root,
            report_path=tmp_path / "historical.csv",
            all_historical=True,
        )
        session.commit()

        assert report.total_files == 1
        assert report.date_confirmed_files == 1
        metadata = session.scalar(select(DocumentMetadataVersion))
        assert metadata is not None
        assert metadata.quote_date.isoformat() == "2024-05-20"


def test_date_extractor_accepts_spaced_label_and_quote_header_date() -> None:
    labeled = _date_text_candidates(
        "견 적 서  제 출 일 : 2021년 07월 15일",
        page=1,
        quote_header=True,
    )
    header = _date_text_candidates(
        "(현대 기아 설비협력업체 견적통일양식) 주식회사 신화 2022-02-24 대표 이사",
        page=1,
        quote_header=True,
    )

    assert {(row.value_text, row.source_kind) for row in labeled} == {
        ("2021-07-15", "EXPLICIT_QUOTE_DATE_TEXT")
    }
    assert {(row.value_text, row.source_kind) for row in header} == {
        ("2022-02-24", "QUOTE_HEADER_DATE")
    }


def test_date_extractor_rejects_quote_number_and_validity_date() -> None:
    candidates = _date_text_candidates(
        "견 적 서 견적번호 TPA20221109-02 견적 유효기간 : 2024-01-31",
        page=1,
        quote_header=True,
    )

    assert candidates == []


def test_date_extractor_reads_english_ocr_header_but_not_validity() -> None:
    candidates = _date_text_candidates(
        "QUOTATION DATE: November 22, 2021 Terms of Validity: November 24, 2021",
        page=1,
        quote_header=True,
        ocr_source=True,
    )

    assert [(row.value_text, row.source_kind, row.confidence) for row in candidates] == [
        ("2021-11-22", "OCR_EXPLICIT_QUOTE_DATE_TEXT", 90)
    ]


def test_date_selection_prefers_cover_date_and_is_field_local() -> None:
    candidates = [
        ExtractedCandidate(
            field_name="quote_date",
            value_text="2024-06-26",
            source_kind="EXPLICIT_QUOTE_DATE_TEXT",
            confidence=99,
            page=1,
        ),
        ExtractedCandidate(
            field_name="quote_date",
            value_text="2024-03-22",
            source_kind="EXPLICIT_QUOTE_DATE_TEXT",
            confidence=99,
            page=3,
        ),
        ExtractedCandidate(
            field_name="supplier_name",
            value_text="업체A",
            source_kind="LABELED_TEXT",
            confidence=95,
            page=1,
        ),
        ExtractedCandidate(
            field_name="supplier_name",
            value_text="업체B",
            source_kind="LABELED_TEXT",
            confidence=95,
            page=1,
        ),
    ]

    selected, ambiguous_fields = _select_candidates(candidates)

    assert selected["quote_date"].value_text == "2024-06-26"
    assert "quote_date" not in ambiguous_fields
    assert "supplier_name" in ambiguous_fields
