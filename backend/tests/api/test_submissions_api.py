from __future__ import annotations

import hashlib
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.catalog.models import (
    ItemMembershipDecision,
    StandardItem,
    StandardPriceVersion,
)
from app.cleansing.models import CleanDecision, CleanStatus
from app.core.config import settings
from app.documents.models import SourceDocument, SourceVariant
from app.quotes.models import RawQuoteItem
from app.standard_database.models import (
    QuoteDocumentPurpose,
    QuoteDocumentRole,
    StandardDatabaseBuildRun,
)
from app.standard_database.service import build_standard_database


def _xlsx_bytes(
    *,
    item_name: str = "Relay",
    spec: str = "24VDC",
    unit: str = "EA",
    price: int = 12000,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "item",
            "spec",
            "unit",
            "quantity",
            "unit price",
            "amount",
            "maker",
        ]
    )
    sheet.append([item_name, spec, unit, 2, price, price * 2, "ACME"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _post(
    client: TestClient,
    content: bytes,
    *,
    filename: str = "new-bid.xlsx",
    submitted_by: str = "buyer-a",
):
    return client.post(
        "/api/submissions",
        data={"submitted_by": submitted_by},
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet",
            )
        },
    )


def _current_role(
    session: Session,
    document_id: int,
) -> QuoteDocumentRole | None:
    return session.scalar(
        select(QuoteDocumentRole)
        .where(QuoteDocumentRole.document_id == document_id)
        .order_by(QuoteDocumentRole.id.desc())
        .limit(1)
    )


def test_upload_ingests_incoming_bid_and_preserves_exact_evidence(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)
    content = _xlsx_bytes()
    digest = hashlib.sha256(content).hexdigest()

    response = _post(client, content)

    assert response.status_code == 201
    payload = response.json()
    assert payload == {
        "document_id": payload["document_id"],
        "sha256": digest,
        "purpose": "INCOMING_BID",
        "parser_name": "quote-reader",
        "parser_version": "reader-v1",
        "status": "INGESTED",
        "raw_item_count": 1,
        "included_count": 1,
        "excluded_count": 0,
        "review_required_count": 0,
    }
    stored = submission_root / digest / "new-bid.xlsx"
    assert stored.read_bytes() == content
    document = api_session.get(SourceDocument, payload["document_id"])
    assert document is not None
    assert _current_role(api_session, document.id).purpose is (
        QuoteDocumentPurpose.INCOMING_BID
    )


def test_upload_rejects_empty_unsupported_and_traversal_filenames(
    client: TestClient,
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)

    empty = _post(client, b"")
    unsupported = _post(client, b"text", filename="quote.txt")
    traversal = _post(client, _xlsx_bytes(), filename="../escape.xlsx")

    assert empty.status_code == 400
    assert empty.json()["detail"]["error_code"] == "EMPTY_UPLOAD"
    assert unsupported.status_code == 415
    assert (
        unsupported.json()["detail"]["error_code"]
        == "UNSUPPORTED_FILE_TYPE"
    )
    assert traversal.status_code == 400
    assert traversal.json()["detail"]["error_code"] == "INVALID_FILENAME"
    assert not (tmp_path / "escape.xlsx").exists()


def test_upload_preserves_valid_special_characters_without_name_collision(
    client: TestClient,
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)
    content = _xlsx_bytes()
    digest = hashlib.sha256(content).hexdigest()
    exact_name = "RFQ#123+가격&조건.xlsx"
    underscore_name = "RFQ_123_가격_조건.xlsx"

    exact = _post(client, content, filename=exact_name)
    underscore = _post(client, content, filename=underscore_name)

    assert exact.status_code == 201
    assert underscore.status_code == 201
    assert exact.json()["document_id"] != underscore.json()["document_id"]
    assert (submission_root / digest / exact_name).read_bytes() == content
    assert (
        submission_root / digest / underscore_name
    ).read_bytes() == content


@pytest.mark.parametrize(
    "filename",
    [
        "CON.xlsx",
        "lpt1.xlsx",
        "trailing-space.xlsx ",
        "trailing-dot.xlsx.",
        "bad\u0001name.xlsx",
        "bad:name.xlsx",
        r"C:\absolute.xlsx",
        r"\\server\share.xlsx",
    ],
)
def test_upload_rejects_unsafe_windows_basenames(
    filename: str,
) -> None:
    from app.api.submissions import _validated_filename

    with pytest.raises(HTTPException) as raised:
        _validated_filename(filename)

    assert raised.value.status_code == 400
    assert raised.value.detail["error_code"] == "INVALID_FILENAME"


def test_upload_rejects_blank_actor_and_configured_size_limit(
    client: TestClient,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "submission_folder", tmp_path / "submitted")
    content = _xlsx_bytes()

    blank_actor = _post(client, content, submitted_by="   ")
    monkeypatch.setattr(settings, "submission_max_bytes", len(content) - 1)
    too_large = _post(client, content, submitted_by="buyer-a")

    assert blank_actor.status_code == 422
    assert (
        blank_actor.json()["detail"]["error_code"]
        == "INVALID_SUBMITTED_BY"
    )
    assert too_large.status_code == 413
    assert too_large.json()["detail"]["error_code"] == "UPLOAD_TOO_LARGE"


def test_malformed_upload_preserves_evidence_without_database_rows(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)
    content = b"not an xlsx archive"
    digest = hashlib.sha256(content).hexdigest()

    response = _post(client, content, filename="broken.xlsx")

    assert response.status_code == 422
    assert response.json()["detail"]["error_code"] == "UNREADABLE_SOURCE"
    assert (submission_root / digest / "broken.xlsx").read_bytes() == content
    assert api_session.scalar(select(func.count(SourceDocument.id))) == 0
    assert api_session.scalar(select(func.count(SourceVariant.id))) == 0
    assert api_session.scalar(select(func.count(RawQuoteItem.id))) == 0


def test_database_failure_rolls_back_ingestion_but_preserves_evidence(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    from app.api import submissions

    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)
    content = _xlsx_bytes()
    digest = hashlib.sha256(content).hexdigest()

    def fail_role_write(*args, **kwargs):
        raise RuntimeError("simulated database write failure")

    monkeypatch.setattr(submissions, "_ensure_incoming_role", fail_role_write)

    with pytest.raises(RuntimeError, match="simulated database"):
        _post(client, content)

    assert (submission_root / digest / "new-bid.xlsx").read_bytes() == content
    assert api_session.scalar(select(func.count(SourceDocument.id))) == 0
    assert api_session.scalar(select(func.count(SourceVariant.id))) == 0
    assert api_session.scalar(select(func.count(RawQuoteItem.id))) == 0
    assert api_session.scalar(select(func.count(CleanDecision.id))) == 0


def test_same_upload_is_idempotent(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "submission_folder", tmp_path / "submitted")
    content = _xlsx_bytes()

    first = _post(client, content)
    second = _post(client, content)

    assert first.status_code == second.status_code == 201
    assert second.json()["document_id"] == first.json()["document_id"]
    assert second.json()["status"] == "UNCHANGED"
    assert api_session.scalar(select(func.count(SourceDocument.id))) == 1
    assert api_session.scalar(select(func.count(SourceVariant.id))) == 1
    assert api_session.scalar(select(func.count(RawQuoteItem.id))) == 1
    assert api_session.scalar(select(func.count(CleanDecision.id))) == 1
    assert api_session.scalar(select(func.count(QuoteDocumentRole.id))) == 1


def test_upload_does_not_override_existing_historical_role(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    submission_root = tmp_path / "submitted"
    monkeypatch.setattr(settings, "submission_folder", submission_root)
    content = _xlsx_bytes()
    digest = hashlib.sha256(content).hexdigest()
    target = submission_root / digest / "new-bid.xlsx"
    target.parent.mkdir(parents=True)
    target.write_bytes(content)
    document = SourceDocument(logical_name=f"{digest}/new-bid.xlsx")
    variant = SourceVariant(
        document=document,
        path=f"{digest}/new-bid.xlsx",
        sha256=digest,
        extension=".xlsx",
        security_state="UNKNOWN",
        selected_for_parsing_at_ingest=True,
    )
    api_session.add(variant)
    api_session.flush()
    api_session.add(
        QuoteDocumentRole(
            document_id=document.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            decided_by="data-owner",
            reason_detail="explicit historical evidence",
        )
    )
    api_session.commit()

    response = _post(client, content)

    assert response.status_code == 409
    assert (
        response.json()["detail"]["error_code"]
        == "DOCUMENT_ROLE_CONFLICT"
    )
    assert _current_role(api_session, document.id).purpose is (
        QuoteDocumentPurpose.HISTORICAL_REFERENCE
    )
    assert api_session.scalar(select(func.count(QuoteDocumentRole.id))) == 1


def test_analysis_detail_accepts_only_current_incoming_bid_role(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "submission_folder", tmp_path / "submitted")
    uploaded = _post(client, _xlsx_bytes()).json()
    incoming = client.get(
        f"/api/analysis/documents/{uploaded['document_id']}"
    )
    historical = SourceDocument(logical_name="historical.xlsx")
    api_session.add(historical)
    api_session.flush()
    api_session.add(
        QuoteDocumentRole(
            document_id=historical.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            decided_by="data-owner",
            reason_detail="training evidence",
        )
    )
    api_session.commit()

    rejected = client.get(f"/api/analysis/documents/{historical.id}")

    assert incoming.status_code == 200
    assert incoming.json()["document"]["purpose"] == "INCOMING_BID"
    assert rejected.status_code == 409
    assert (
        rejected.json()["detail"]["error_code"]
        == "DOCUMENT_ROLE_MISMATCH"
    )


def test_upload_creates_no_catalog_memberships_or_prices(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "submission_folder", tmp_path / "submitted")

    response = _post(client, _xlsx_bytes())

    assert response.status_code == 201
    assert api_session.scalar(select(func.count(StandardItem.id))) == 0
    assert (
        api_session.scalar(select(func.count(ItemMembershipDecision.id))) == 0
    )
    assert (
        api_session.scalar(select(func.count(StandardPriceVersion.id))) == 0
    )


def test_upload_does_not_change_standard_build_input_or_create_new_run(
    client: TestClient,
    api_session: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    historical = SourceDocument(logical_name="historical.xlsx")
    historical_variant = SourceVariant(
        document=historical,
        path="historical.xlsx",
        sha256="a" * 64,
        extension=".xlsx",
        security_state="UNKNOWN",
        selected_for_parsing_at_ingest=True,
    )
    raw = RawQuoteItem(
        source_variant=historical_variant,
        source_sheet="Sheet1",
        source_row=2,
        item_name_raw="Relay",
        spec_raw="24VDC",
        unit_raw="EA",
        unit_price_raw="12000",
        parser_name="quote-reader",
        parser_version="reader-v1",
    )
    api_session.add(
        CleanDecision(
            raw_item=raw,
            status=CleanStatus.INCLUDED,
            reason_code="VALID",
            item_name_norm="RELAY",
            spec_norm="24VDC",
            unit_norm="EA",
            unit_price=Decimal("12000"),
            rule_version="clean-v1",
        )
    )
    api_session.flush()
    api_session.add(
        QuoteDocumentRole(
            document_id=historical.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            decided_by="data-owner",
            reason_detail="training evidence",
        )
    )
    first = build_standard_database(api_session)
    api_session.commit()
    first_run = api_session.get(StandardDatabaseBuildRun, first.run_id)
    assert first_run is not None
    first_fingerprint = first_run.input_fingerprint
    monkeypatch.setattr(settings, "submission_folder", tmp_path / "submitted")

    uploaded = _post(client, _xlsx_bytes(price=999999))
    second = build_standard_database(api_session)

    assert uploaded.status_code == 201
    assert second.reused_run_id == first.run_id
    assert api_session.get(
        StandardDatabaseBuildRun, second.run_id
    ).input_fingerprint == first_fingerprint
    assert (
        api_session.scalar(select(func.count(StandardDatabaseBuildRun.id)))
        == 1
    )
