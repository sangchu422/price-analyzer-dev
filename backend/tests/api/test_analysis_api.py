from __future__ import annotations

from datetime import date
from decimal import Decimal

import app.analysis.target_price as target_price
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.catalog.models import (
    DocumentMetadataVersion,
    ItemMembershipDecision,
    StandardItem,
    StandardPriceVersion,
)
from app.analysis.models import (
    InflationIndexPoint,
    InflationSyncRun,
    QuoteAnalysisLineResult,
    QuoteAnalysisRun,
)
from app.cleansing.models import CleanDecision, CleanStatus
from app.documents.models import SourceDocument, SourceVariant
from app.quotes.models import RawQuoteItem
from app.procurement.models import (
    ItemCategory,
    QuoteCatalogActivationEntry,
    QuoteCatalogActivationRun,
    StandardItemCategoryAssignment,
)
from app.standard_database.models import (
    QuoteDocumentPurpose,
    QuoteDocumentRole,
)
from app.standard_database.service import build_standard_database


def _document(
    session: Session,
    *,
    name: str = "quotes/new.xlsx",
    rows: int = 2,
    duplicate_items: bool = False,
) -> SourceDocument:
    document = SourceDocument(logical_name=name)
    variant = SourceVariant(
        document=document,
        path=name,
        sha256="a" * 64,
        extension=".xlsx",
        security_state="UNLOCKED",
        selected_for_parsing_at_ingest=True,
    )
    for row in range(1, rows + 1):
        item_number = 1 if duplicate_items else row
        raw = RawQuoteItem(
            source_variant=variant,
            source_sheet="Sheet1",
            source_row=row,
            source_cells=f"A{row}:G{row}",
            item_name_raw=f"CUSTOM ITEM {item_number}",
            spec_raw=f"ZZ-{item_number}",
            unit_raw="EA",
            unit_price_raw=str(row * 100),
            parser_name="xlsx",
            parser_version="reader-v1",
        )
        session.add(
            CleanDecision(
                raw_item=raw,
                status=CleanStatus.INCLUDED,
                reason_code="VALID",
                item_name_norm=f"CUSTOM ITEM {item_number}",
                spec_norm=f"ZZ-{item_number}",
                unit_norm="EA",
                unit_price=Decimal(row * 100),
                rule_version="clean-v1",
            )
        )
    session.add(document)
    session.flush()
    session.add(
        QuoteDocumentRole(
            document_id=document.id,
            purpose=QuoteDocumentPurpose.INCOMING_BID,
            decided_by="test-submitter",
            reason_detail="API analysis fixture",
        )
    )
    session.commit()
    return document


def test_analysis_document_list_and_typed_detail(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session)

    listing = client.get("/api/analysis/documents?limit=10&offset=0")
    detail = client.get(
        f"/api/analysis/documents/{document.id}?limit=1"
    )

    assert listing.status_code == 200
    assert listing.json()["items"] == [
        {
            "id": document.id,
            "logical_name": "quotes/new.xlsx",
            "raw_item_count": 2,
            "included_count": 2,
            "excluded_count": 0,
            "review_required_count": 0,
            "undecided_count": 0,
            "analysis_ready": True,
        }
    ]
    assert listing.json()["total"] == 1
    assert detail.status_code == 200
    payload = detail.json()
    assert payload["document"]["id"] == document.id
    assert payload["document"]["display_name"] == "new.xlsx"
    assert payload["price_policy"] == {
        "within_percent": "10",
        "high_low_percent": "20",
        "description": (
            "표준 중앙값 대비 ±10% 이내 적정, ±10~20% 주의, "
            "±20% 초과 고가·저가"
        ),
    }
    assert len(payload["lines"]) == 1
    assert payload["lines"][0]["match_status"] == "NO_MATCH"
    assert payload["lines"][0]["canonical_name"] is None
    assert payload["lines"][0]["canonical_spec"] is None
    assert payload["lines"][0]["canonical_unit"] is None
    assert payload["lines"][0]["source"]["path"] == "quotes/new.xlsx"
    assert payload["next_cursor"] == payload["lines"][0]["raw_item_id"]


def test_analysis_run_persists_thresholds_and_keeps_market_target_separate(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, rows=1)

    response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={
            "created_by": "buyer-01",
            "review_percent": 12,
            "high_percent": 25,
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["price_policy"]["within_percent"] == "12"
    assert payload["price_policy"]["high_low_percent"] == "25"
    assert payload["target_period"] is None
    assert payload["target_lines"][0]["status"] == "MARKET_REFERENCE_REQUIRED"
    assert payload["target_lines"][0]["target_unit_price"] is None
    assert payload["family_analysis"]["rule_version"] == "item-family-keyword-v1"
    assert payload["family_analysis"]["matched_count"] == 0
    assert payload["family_analysis"]["pending_count"] == 1
    run = api_session.get(QuoteAnalysisRun, payload["run_id"])
    assert run is not None
    assert run.review_percent == Decimal("12")
    assert run.high_percent == Decimal("25")
    stored = api_session.scalar(
        select(QuoteAnalysisLineResult).where(
            QuoteAnalysisLineResult.analysis_run_id == run.id
        )
    )
    assert stored is not None
    assert stored.target_status == "MARKET_REFERENCE_REQUIRED"
    saved = client.get(f"/api/analysis/runs/{run.id}")
    assert saved.status_code == 200
    assert saved.json()["review_percent"] == "12.000000"
    assert saved.json()["target_lines"][0]["status"] == "MARKET_REFERENCE_REQUIRED"


def test_analysis_run_rejects_reversed_thresholds(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, rows=1)

    response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={
            "created_by": "buyer-01",
            "review_percent": 30,
            "high_percent": 20,
        },
    )

    assert response.status_code == 422


def test_activation_validates_outlook_before_catalog_commit(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, rows=1)
    run_response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={
            "created_by": "buyer-01",
            "review_percent": 10,
            "high_percent": 20,
        },
    )
    run_id = run_response.json()["run_id"]

    response = client.post(
        f"/api/analysis/runs/{run_id}/activate",
        json={
            "activated_by": "buyer-01",
            "reason_detail": "검토 완료 후 신규 견적 반영",
            "send_outlook": True,
            "outlook_recipient": "invalid-address",
        },
    )

    assert response.status_code == 422
    assert "이메일" in response.json()["detail"]
    assert api_session.scalar(
        select(func.count(QuoteCatalogActivationRun.id))
    ) == 0
    assert api_session.scalar(
        select(QuoteDocumentRole.purpose)
        .where(QuoteDocumentRole.document_id == document.id)
        .order_by(QuoteDocumentRole.id.desc())
        .limit(1)
    ) == QuoteDocumentPurpose.INCOMING_BID


def test_activation_appends_new_catalog_evidence_once(
    client: TestClient,
    api_session: Session,
) -> None:
    api_session.add(
        ItemCategory(
            code="GENERAL_COMPONENT",
            name="공통 설비·부품",
            description="테스트용 기본 분류",
            sort_order=900,
        )
    )
    api_session.commit()
    document = _document(api_session, rows=2, duplicate_items=True)
    run_response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={
            "created_by": "buyer-01",
            "review_percent": 10,
            "high_percent": 20,
        },
    )
    assert run_response.status_code == 200, run_response.text
    run_id = run_response.json()["run_id"]

    first = client.post(
        f"/api/analysis/runs/{run_id}/activate",
        json={
            "activated_by": "buyer-01",
            "reason_detail": "검토 완료 후 신규 견적 반영",
            "send_outlook": False,
        },
    )

    assert first.status_code == 200, first.text
    payload = first.json()
    assert payload["status"] == "SUCCEEDED"
    assert payload["counts"]["included_rows"] == 2
    assert payload["counts"]["created_standard_items"] == 1
    assert payload["counts"]["matched_activation_item_rows"] == 1
    assert payload["counts"]["price_versions_created"] == 1
    assert len(payload["entries"]) == 2
    assert all(row["standard_price_version_id"] for row in payload["entries"])
    assert len({row["standard_item_id"] for row in payload["entries"]}) == 1
    assert api_session.scalar(select(func.count(StandardItem.id))) == 1
    assert api_session.scalar(select(func.count(StandardPriceVersion.id))) == 1
    assert api_session.scalar(
        select(func.count(StandardItemCategoryAssignment.id))
    ) == 1
    assert api_session.scalar(
        select(QuoteDocumentRole.purpose)
        .where(QuoteDocumentRole.document_id == document.id)
        .order_by(QuoteDocumentRole.id.desc())
        .limit(1)
    ) == QuoteDocumentPurpose.HISTORICAL_REFERENCE

    second = client.post(
        f"/api/analysis/runs/{run_id}/activate",
        json={
            "activated_by": "buyer-01",
            "reason_detail": "동일 요청 재시도",
            "send_outlook": False,
        },
    )
    assert second.status_code == 200, second.text
    assert second.json()["activation_run_id"] == payload["activation_run_id"]
    assert api_session.scalar(
        select(func.count(QuoteCatalogActivationRun.id))
    ) == 1
    assert api_session.scalar(
        select(func.count(QuoteCatalogActivationEntry.id))
    ) == 2
    assert api_session.scalar(select(func.count(StandardPriceVersion.id))) == 1


def test_incoming_exact_key_uses_standard_price_without_membership_write(
    client: TestClient,
    api_session: Session,
) -> None:
    historical = SourceDocument(logical_name="historical.xlsx")
    variant = SourceVariant(
        document=historical,
        path="historical.xlsx",
        sha256="b" * 64,
        extension=".xlsx",
        security_state="UNLOCKED",
        selected_for_parsing_at_ingest=True,
    )
    raw = RawQuoteItem(
        source_variant=variant,
        source_sheet="Sheet1",
        source_row=1,
        item_name_raw="CUSTOM ITEM 1",
        spec_raw="ZZ-1",
        unit_raw="EA",
        quantity_raw="1",
        unit_price_raw="80",
        amount_raw="80",
        parser_name="xlsx",
        parser_version="reader-v1",
    )
    api_session.add(
        CleanDecision(
            raw_item=raw,
            status=CleanStatus.INCLUDED,
            reason_code="VALID",
            item_name_norm="CUSTOM ITEM 1",
            spec_norm="ZZ-1",
            unit_norm="EA",
            quantity=Decimal("1"),
            unit_price=Decimal("80"),
            amount=Decimal("80"),
            rule_version="clean-v1",
        )
    )
    api_session.flush()
    api_session.add(
        QuoteDocumentRole(
            document_id=historical.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            decided_by="data-owner",
            reason_detail="training evidence",
        )
    )
    build_standard_database(api_session)
    api_session.commit()
    incoming = _document(api_session, rows=1)
    before_memberships = api_session.scalar(
        select(func.count(ItemMembershipDecision.id))
    )
    before_prices = api_session.scalar(
        select(func.count(StandardPriceVersion.id))
    )

    response = client.get(
        f"/api/analysis/documents/{incoming.id}?limit=100"
    )

    assert response.status_code == 200
    line = response.json()["lines"][0]
    assert line["match_status"] == "MATCHED"
    assert line["membership_decision_id"] is None
    assert line["quantity"] is None
    assert line["quote_amount"] is None
    assert api_session.scalar(
        select(func.count(ItemMembershipDecision.id))
    ) == before_memberships
    assert api_session.scalar(
        select(func.count(StandardPriceVersion.id))
    ) == before_prices


def test_analysis_list_excludes_historical_and_unclassified_documents(
    client: TestClient,
    api_session: Session,
) -> None:
    incoming = _document(api_session, name="incoming.xlsx", rows=1)
    historical = SourceDocument(logical_name="historical.xlsx")
    unclassified = SourceDocument(logical_name="unclassified.xlsx")
    api_session.add_all([historical, unclassified])
    api_session.flush()
    api_session.add(
        QuoteDocumentRole(
            document_id=historical.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            decided_by="data-owner",
            reason_detail="training evidence",
        )
    )
    api_session.commit()

    response = client.get("/api/analysis/documents?limit=10")

    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert [item["id"] for item in response.json()["items"]] == [incoming.id]


def test_analysis_list_uses_only_latest_document_role(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, name="reclassified.xlsx", rows=1)
    current = api_session.scalar(
        select(QuoteDocumentRole)
        .where(QuoteDocumentRole.document_id == document.id)
        .order_by(QuoteDocumentRole.id.desc())
    )
    api_session.add(
        QuoteDocumentRole(
            document_id=document.id,
            purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
            supersedes_role_id=current.id,
            decided_by="data-owner",
            reason_detail="corrected classification",
        )
    )
    api_session.commit()

    response = client.get("/api/analysis/documents?limit=10")

    assert response.status_code == 200
    assert response.json()["total"] == 0
    assert response.json()["items"] == []


def test_detail_cursor_and_status_filter_are_stable(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, rows=3)

    first = client.get(
        f"/api/analysis/documents/{document.id}"
        "?limit=1&match_status=NO_MATCH"
    ).json()
    second = client.get(
        f"/api/analysis/documents/{document.id}"
        f"?limit=1&match_status=NO_MATCH&after_id={first['next_cursor']}"
    ).json()

    assert first["lines"][0]["raw_item_id"] < (
        second["lines"][0]["raw_item_id"]
    )
    assert first["next_cursor"] is not None


def test_refresh_candidates_never_creates_membership(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session)
    before = api_session.scalar(
        select(func.count(ItemMembershipDecision.id))
    )

    response = client.post(
        f"/api/analysis/documents/{document.id}/refresh-candidates"
    )

    assert response.status_code == 200
    assert response.json()["refreshed_candidate_rows"] == 2
    assert response.json()["membership_rows_created"] == 0
    assert (
        api_session.scalar(select(func.count(ItemMembershipDecision.id)))
        == before
    )


def test_analysis_api_rejects_missing_documents_and_bad_page_bounds(
    client: TestClient,
) -> None:
    assert client.get("/api/analysis/documents/999").status_code == 404
    assert (
        client.get("/api/analysis/documents/999?limit=101").status_code
        == 422
    )
    assert (
        client.get("/api/analysis/documents?offset=-1").status_code
        == 422
    )


def test_document_analysis_export_returns_xlsx_workbook(
    client: TestClient,
    api_session: Session,
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    document = _document(api_session, rows=2)

    response = client.get(
        f"/api/analysis/documents/{document.id}/export"
        "?review_percent=15&high_percent=25"
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "품명", "규격", "단위", "수량", "개당 단가", "구매 금액",
        "참조 기준가", "참조 최저", "참조 최고", "편차 금액", "편차율(%)",
        "매칭 상태", "표준 품목 ID", "표준 가격 버전 ID", "가격 판정",
    )
    assert len(rows) == 3
    assert {row[0] for row in rows[1:]} == {"CUSTOM ITEM 1", "CUSTOM ITEM 2"}


def test_cpi_sync_api_returns_confirmed_annual_rate_evidence(
    client: TestClient,
    monkeypatch,
) -> None:
    rates = {
        "2017": "1.9",
        "2018": "1.5",
        "2019": "0.4",
        "2020": "0.5",
        "2021": "2.5",
        "2022": "5.1",
        "2023": "3.6",
        "2024": "2.3",
        "2025": "2.1",
    }
    payload = [
        {
            "ORG_ID": "101",
            "TBL_ID": "DT_1J22041",
            "ITM_ID": "T",
            "C1": "0",
            "PRD_SE": "A",
            "PRD_DE": year,
            "DT": rate,
            "LST_CHN_DE": "2026-01-15",
        }
        for year, rate in rates.items()
    ]
    request_params: dict[str, object] = {}

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> list[dict[str, str]]:
            return payload

    def fake_get(url: str, **kwargs: object) -> FakeResponse:
        assert url.endswith("/v1/kosis/data")
        request_params.update(kwargs["params"])
        assert kwargs["headers"] == {"User-Agent": "price-analyzer/cpi-sync"}
        return FakeResponse()

    monkeypatch.setattr(target_price.httpx, "get", fake_get)

    response = client.post("/api/analysis/inflation/series/cpi-all/sync")

    assert response.status_code == 200, response.text
    assert request_params["orgId"] == "101"
    assert request_params["tblId"] == "DT_1J22041"
    assert request_params["itmId"] == "T"
    assert request_params["objL1"] == "0"
    assert request_params["prdSe"] == "Y"
    body = response.json()
    assert body["available"] is True
    assert body["latest_period"] == "2025"
    assert body["latest_annual_rate"] == "2.100000"
    assert body["factor"] == "1.216544"
    assert body["cumulative_percent"] == "21.654370"
    assert body["source_url"] == (
        "https://kosis.kr/statHtml/statHtml.do?orgId=101&tblId=DT_1J22041"
    )
    assert body["sync_run_id"] is not None
    assert body["annual_rates"] == [
        {"year": year, "rate": f"{Decimal(rate):.6f}"}
        for year, rate in rates.items()
    ]

    cached = client.get("/api/analysis/inflation/series/cpi-all")
    assert cached.status_code == 200
    assert cached.json()["sync_run_id"] == body["sync_run_id"]
    assert cached.json()["latest_period"] == "2025"

    legacy_ppi = client.get("/api/analysis/inflation/series/ppi-all")
    assert legacy_ppi.status_code == 200
    assert legacy_ppi.json()["available"] is False


def test_cpi_and_legacy_ppi_caches_remain_separate(
    client: TestClient,
    api_session: Session,
) -> None:
    ppi = InflationSyncRun(
        series_kind="PPI_ALL",
        provider="KOSIS",
        org_id="301",
        table_id="DT_404Y014",
        item_id="13103134604999",
        classifier_code="13102134604ACC_CD.*AA",
        period_type="M",
        unit="2020=100",
        source_url="https://example.test/ppi",
        source_last_changed=None,
        response_sha256="p" * 64,
        row_count=1,
        latest_period="202606",
    )
    api_session.add(ppi)
    api_session.flush()
    api_session.add(
        InflationIndexPoint(
            sync_run_id=ppi.id,
            period="202606",
            index_value=Decimal("130.03"),
        )
    )
    cpi = InflationSyncRun(
        series_kind="CPI_ALL",
        provider="KOSIS",
        org_id="101",
        table_id="DT_1J22041",
        item_id="T",
        classifier_code="0",
        period_type="Y",
        unit="%",
        source_url="https://example.test/cpi",
        source_last_changed=None,
        response_sha256="c" * 64,
        row_count=2,
        latest_period="2025",
    )
    api_session.add(cpi)
    api_session.flush()
    api_session.add_all(
        [
            InflationIndexPoint(
                sync_run_id=cpi.id,
                period="2024",
                index_value=Decimal("2.3"),
            ),
            InflationIndexPoint(
                sync_run_id=cpi.id,
                period="2025",
                index_value=Decimal("2.1"),
            ),
        ]
    )
    api_session.commit()

    ppi_response = client.get("/api/analysis/inflation/series/ppi-all")
    cpi_response = client.get("/api/analysis/inflation/series/cpi-all")

    assert ppi_response.status_code == 200
    assert ppi_response.json()["sync_run_id"] == ppi.id
    assert ppi_response.json()["latest_period"] == "202606"
    assert cpi_response.status_code == 200
    assert cpi_response.json()["sync_run_id"] == cpi.id
    assert cpi_response.json()["latest_period"] == "2025"


def test_new_analysis_run_uses_cpi_and_never_falls_back_to_ppi(
    client: TestClient,
    api_session: Session,
) -> None:
    document = _document(api_session, rows=1)
    cpi = InflationSyncRun(
        series_kind="CPI_ALL",
        provider="KOSIS",
        org_id="101",
        table_id="DT_1J22041",
        item_id="T",
        classifier_code="0",
        period_type="Y",
        unit="%",
        source_url="https://kosis.kr/statHtml/statHtml.do?orgId=101&tblId=DT_1J22041",
        source_last_changed=None,
        response_sha256="a" * 64,
        row_count=1,
        latest_period="2025",
    )
    api_session.add(cpi)
    api_session.flush()
    api_session.add(
        InflationIndexPoint(
            sync_run_id=cpi.id,
            period="2025",
            index_value=Decimal("2.1"),
        )
    )
    # This newer PPI snapshot must not alter a new analysis run's CPI basis.
    ppi = InflationSyncRun(
        series_kind="PPI_ALL",
        provider="KOSIS",
        org_id="301",
        table_id="DT_404Y014",
        item_id="13103134604999",
        classifier_code="13102134604ACC_CD.*AA",
        period_type="M",
        unit="2020=100",
        source_url="https://example.test/ppi",
        source_last_changed=None,
        response_sha256="b" * 64,
        row_count=1,
        latest_period="202606",
    )
    api_session.add(ppi)
    api_session.flush()
    api_session.add(
        InflationIndexPoint(
            sync_run_id=ppi.id,
            period="202606",
            index_value=Decimal("130.03"),
        )
    )
    api_session.commit()

    response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={"created_by": "buyer-01"},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["inflation_sync_run_id"] == cpi.id
    assert body["inflation_series_kind"] == "CPI_ALL"
    assert body["target_period"] == "2025"
    assert body["inflation_source_url"] == cpi.source_url


def test_target_price_export_reads_stored_run_without_side_effects(
    client: TestClient,
    api_session: Session,
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    document = _document(api_session, rows=2)
    run_response = client.post(
        f"/api/analysis/documents/{document.id}/runs",
        json={"created_by": "buyer-01", "review_percent": 10, "high_percent": 20},
    )
    assert run_response.status_code == 200, run_response.text
    run_id = run_response.json()["run_id"]

    run_count_before = api_session.scalar(
        select(func.count(QuoteAnalysisRun.id))
    )

    response = client.get(f"/api/analysis/runs/{run_id}/target-price-export")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "품명", "규격", "단위", "수량", "개당 단가", "구매 금액",
        "구매 목표 단가(개당)", "구매 목표금액", "목표 인하 금액", "산정 상태",
    )
    assert len(rows) == 4
    assert {row[0] for row in rows[1:-1]} == {"CUSTOM ITEM 1", "CUSTOM ITEM 2"}
    assert rows[-1][0] == "합계"

    api_session.expire_all()
    run_count_after = api_session.scalar(select(func.count(QuoteAnalysisRun.id)))
    assert run_count_after == run_count_before


def test_target_price_export_rejects_missing_run(client: TestClient) -> None:
    assert (
        client.get("/api/analysis/runs/999/target-price-export").status_code == 404
    )


def test_matched_line_evidence_quality_reflects_distinct_suppliers(
    client: TestClient,
    api_session: Session,
) -> None:
    for row, supplier in [(1, "SUPPLIER Z"), (2, "SUPPLIER Z")]:
        historical = SourceDocument(logical_name=f"historical-{row}.xlsx")
        variant = SourceVariant(
            document=historical,
            path=f"historical-{row}.xlsx",
            sha256=f"{row:064x}",
            extension=".xlsx",
            security_state="UNLOCKED",
            selected_for_parsing_at_ingest=True,
        )
        raw = RawQuoteItem(
            source_variant=variant,
            source_sheet="Sheet1",
            source_row=1,
            item_name_raw="CUSTOM ITEM 1",
            spec_raw="ZZ-1",
            unit_raw="EA",
            unit_price_raw="80",
            parser_name="xlsx",
            parser_version="reader-v1",
        )
        api_session.add_all(
            [
                CleanDecision(
                    raw_item=raw,
                    status=CleanStatus.INCLUDED,
                    reason_code="VALID",
                    item_name_norm="CUSTOM ITEM 1",
                    spec_norm="ZZ-1",
                    unit_norm="EA",
                    unit_price=Decimal("80"),
                    rule_version="clean-v1",
                ),
                DocumentMetadataVersion(
                    source_document=historical,
                    version_number=1,
                    supplier_name=supplier,
                    quote_date=date(2026, 7, row),
                    project_name=None,
                    decided_by="data-owner",
                ),
            ]
        )
        api_session.flush()
        api_session.add(
            QuoteDocumentRole(
                document_id=historical.id,
                purpose=QuoteDocumentPurpose.HISTORICAL_REFERENCE,
                decided_by="data-owner",
                reason_detail="training evidence",
            )
        )
    build_standard_database(api_session)
    api_session.commit()
    incoming = _document(api_session, rows=1)

    response = client.get(
        f"/api/analysis/documents/{incoming.id}?limit=100"
    )

    assert response.status_code == 200
    line = response.json()["lines"][0]
    assert line["match_status"] == "MATCHED"
    assert line["standard_observation_count"] == 2
    assert line["evidence_quality"] == "SINGLE_OBSERVATION"
