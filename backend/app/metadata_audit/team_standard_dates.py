"""Backfill quote dates from the team's legacy standard-price workbook.

The workbook is never treated as a price source.  Only its raw-data date
column is reconciled to an already-ingested historical source document.
"""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.catalog.models import DocumentMetadataVersion
from app.catalog.service import current_document_metadata
from app.documents.models import SourceVariant
from app.standard_database.models import (
    QuoteDocumentPurpose,
    QuoteDocumentRole,
)

BACKFILL_ACTOR = "team-standard-date-backfill-v1"
PLACEHOLDER_DATES = {"2025-01-01"}
EXPECTED_STANDARD_HEADERS = {
    "A3": "표준품목ID",
    "K3": "최근견적일",
}
EXPECTED_RAW_HEADERS = {
    "A1": "ID",
    "C1": "견적번호",
    "D1": "견적일",
    "H1": "품명",
    "I1": "규격",
    "J1": "단위",
    "K1": "수량",
    "L1": "단가(원)",
    "M1": "금액(원)",
}


class TeamStandardDateError(ValueError):
    """The workbook, companion JSON, or reconciliation input is invalid."""


@dataclass(frozen=True)
class TeamStandardDateReport:
    workbook_sha256: str
    source_json_sha256: str
    source_revision: str
    workbook_raw_rows: int
    workbook_standard_rows: int
    workbook_standard_valid_dates: int
    workbook_standard_invalid_dates: int
    source_files: int
    matched_files: int
    unmatched_files: int
    ambiguous_files: int
    placeholder_files: int
    no_date_files: int
    eligible_files: int
    applied_files: int
    already_same_files: int
    conflict_files: int
    manual_reference_files: int
    inferred_file_date_files: int
    applied_document_ids: tuple[int, ...]
    conflicts: tuple[dict[str, Any], ...]
    unmatched: tuple[str, ...]
    ambiguous: tuple[dict[str, Any], ...]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class _SourceDate:
    source_path: str
    quote_date: date | None
    quote_number: str | None
    source_kind: str
    quality: str
    confidence: int
    use_for_index: str
    placeholder: bool


def backfill_team_standard_dates(
    session: Session,
    *,
    workbook_path: Path,
    source_json_path: Path,
    source_revision: str,
    apply: bool,
) -> TeamStandardDateReport:
    workbook_path = workbook_path.resolve(strict=True)
    source_json_path = source_json_path.resolve(strict=True)
    workbook_sha = _sha256(workbook_path)
    source_json_sha = _sha256(source_json_path)
    workbook_stats = _validate_workbook(workbook_path, source_json_path)
    source_dates = _load_source_dates(source_json_path)
    variants = _historical_variants(session)
    matched, unmatched, ambiguous = _match_sources(source_dates, variants)

    applied_document_ids: list[int] = []
    conflicts: list[dict[str, Any]] = []
    already_same = 0
    placeholder = 0
    no_date = 0
    eligible = 0
    manual_reference = 0
    inferred_file_date = 0

    for source, variant in matched:
        if source.placeholder:
            placeholder += 1
            continue
        if source.quote_date is None:
            no_date += 1
            continue
        eligible += 1
        if source.quality == "REFERENCE_BACKFILL":
            manual_reference += 1
        if source.quality == "FILE_DATE_INFERRED":
            inferred_file_date += 1

        current = current_document_metadata(session, variant.document_id)
        if current is not None and current.quote_date is not None:
            if current.quote_date == source.quote_date:
                already_same += 1
            else:
                conflicts.append({
                    "source_path": source.source_path,
                    "document_id": variant.document_id,
                    "current_quote_date": current.quote_date.isoformat(),
                    "team_quote_date": source.quote_date.isoformat(),
                    "resolution": "KEPT_CURRENT_SOURCE_DATE",
                })
            continue

        if not apply:
            applied_document_ids.append(variant.document_id)
            continue

        prior_evidence = _evidence_dict(current)
        prior_evidence["quote_date"] = {
            "source_kind": source.source_kind,
            "quality": source.quality,
            "confidence": source.confidence,
            "use_for_index": source.use_for_index,
            "workbook_sha256": workbook_sha,
            "source_json_sha256": source_json_sha,
            "source_revision": source_revision,
            "workbook_sheet": "원본견적데이터",
            "workbook_column": "D",
            "legacy_summary_sheet": "표준단가DB",
            "legacy_summary_column": "K",
            "source_path": source.source_path,
            "quote_number": source.quote_number,
            "note": (
                "K열 집계값은 사용하지 않고 동일 원본 경로의 D열 날짜를 "
                "사용함"
            ),
        }
        session.add(
            DocumentMetadataVersion(
                source_document_id=variant.document_id,
                version_number=1 if current is None else current.version_number + 1,
                supplier_name=(
                    None if current is None else current.supplier_name
                ),
                quote_date=source.quote_date,
                project_name=None if current is None else current.project_name,
                decided_by=BACKFILL_ACTOR,
                reason_detail=(
                    "팀 표준단가DB 원본행과 현재 원본 파일을 대조해 견적일 보완"
                ),
                evidence_json=json.dumps(
                    prior_evidence,
                    ensure_ascii=False,
                    separators=(",", ":"),
                ),
            )
        )
        session.flush()
        applied_document_ids.append(variant.document_id)

    return TeamStandardDateReport(
        workbook_sha256=workbook_sha,
        source_json_sha256=source_json_sha,
        source_revision=source_revision,
        workbook_raw_rows=workbook_stats["raw_rows"],
        workbook_standard_rows=workbook_stats["standard_rows"],
        workbook_standard_valid_dates=workbook_stats["standard_valid_dates"],
        workbook_standard_invalid_dates=workbook_stats["standard_invalid_dates"],
        source_files=len(source_dates),
        matched_files=len(matched),
        unmatched_files=len(unmatched),
        ambiguous_files=len(ambiguous),
        placeholder_files=placeholder,
        no_date_files=no_date,
        eligible_files=eligible,
        applied_files=len(applied_document_ids),
        already_same_files=already_same,
        conflict_files=len(conflicts),
        manual_reference_files=manual_reference,
        inferred_file_date_files=inferred_file_date,
        applied_document_ids=tuple(applied_document_ids),
        conflicts=tuple(conflicts),
        unmatched=tuple(unmatched),
        ambiguous=tuple(ambiguous),
    )


def _validate_workbook(
    workbook_path: Path,
    source_json_path: Path,
) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if not {"표준단가DB", "원본견적데이터"}.issubset(workbook.sheetnames):
            raise TeamStandardDateError("required workbook sheets are missing")
        standard = workbook["표준단가DB"]
        raw = workbook["원본견적데이터"]
        _validate_headers(standard, EXPECTED_STANDARD_HEADERS)
        _validate_headers(raw, EXPECTED_RAW_HEADERS)
        json_rows = json.loads(source_json_path.read_text(encoding="utf-8"))
        raw_rows = list(
            raw.iter_rows(min_row=2, max_col=13, values_only=True)
        )
        if len(raw_rows) != len(json_rows):
            raise TeamStandardDateError(
                "workbook raw row count does not match companion JSON"
            )
        for excel_row, json_row in zip(raw_rows, json_rows, strict=True):
            comparisons = (
                (excel_row[0], json_row.get("ID")),
                (excel_row[2], json_row.get("견적번호")),
                (excel_row[3], json_row.get("견적일")),
                (excel_row[7], json_row.get("품명")),
                (excel_row[8], json_row.get("규격")),
                (excel_row[9], json_row.get("단위")),
            )
            if any(_cell_text(left) != _cell_text(right) for left, right in comparisons):
                raise TeamStandardDateError(
                    f"workbook and companion JSON differ at raw ID {excel_row[0]}"
                )
        standard_dates = [
            _cell_text(row[0])
            for row in standard.iter_rows(
                min_row=4,
                min_col=11,
                max_col=11,
                values_only=True,
            )
        ]
        valid = sum(_parse_date(value) is not None for value in standard_dates)
        return {
            "raw_rows": len(raw_rows),
            "standard_rows": len(standard_dates),
            "standard_valid_dates": valid,
            "standard_invalid_dates": len(standard_dates) - valid,
        }
    finally:
        workbook.close()


def _validate_headers(sheet, expected: dict[str, str]) -> None:
    for address, value in expected.items():
        if _cell_text(sheet[address].value) != value:
            raise TeamStandardDateError(
                f"unexpected header at {sheet.title}!{address}"
            )


def _load_source_dates(source_json_path: Path) -> dict[str, _SourceDate]:
    rows = json.loads(source_json_path.read_text(encoding="utf-8"))
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[_cell_text(row.get("출처파일"))].append(row)
    result: dict[str, _SourceDate] = {}
    for source_path, source_rows in grouped.items():
        raw_dates = {_cell_text(row.get("견적일")) for row in source_rows}
        quote_numbers = {
            _cell_text(row.get("견적번호")) for row in source_rows
        }
        if len(raw_dates) != 1 or len(quote_numbers) != 1:
            raise TeamStandardDateError(
                f"source file has conflicting metadata: {source_path}"
            )
        raw_date = next(iter(raw_dates))
        quote_date = _parse_date(raw_date)
        if "바츠 추출 견적서" in source_path:
            source_kind = "TEAM_FILE_TIMESTAMP"
            quality = "FILE_DATE_INFERRED"
            confidence = 40
            use_for_index = "YEAR_ONLY"
        elif "AONE 추출 견적서" in source_path:
            source_kind = "TEAM_AONE_UNKNOWN"
            quality = "UNKNOWN"
            confidence = 0
            use_for_index = "DO_NOT_USE"
        else:
            source_kind = "TEAM_MANUAL_METADATA"
            quality = "REFERENCE_BACKFILL"
            confidence = 70
            use_for_index = "EXACT_DATE"
        result[source_path] = _SourceDate(
            source_path=source_path,
            quote_date=quote_date,
            quote_number=next(iter(quote_numbers)) or None,
            source_kind=source_kind,
            quality=quality,
            confidence=confidence,
            use_for_index=use_for_index,
            placeholder=raw_date in PLACEHOLDER_DATES,
        )
    return result


def _historical_variants(session: Session) -> list[SourceVariant]:
    latest_roles = (
        select(
            QuoteDocumentRole.document_id.label("document_id"),
            func.max(QuoteDocumentRole.id).label("role_id"),
        )
        .group_by(QuoteDocumentRole.document_id)
        .subquery()
    )
    current_roles = {
        document_id: purpose
        for document_id, purpose in session.execute(
            select(
                QuoteDocumentRole.document_id,
                QuoteDocumentRole.purpose,
            ).join(
                latest_roles,
                latest_roles.c.role_id == QuoteDocumentRole.id,
            )
        )
    }
    return [
        variant
        for variant in session.scalars(select(SourceVariant))
        if current_roles.get(variant.document_id)
        == QuoteDocumentPurpose.HISTORICAL_REFERENCE
    ]


def _match_sources(
    source_dates: dict[str, _SourceDate],
    variants: list[SourceVariant],
) -> tuple[
    list[tuple[_SourceDate, SourceVariant]],
    list[str],
    list[dict[str, Any]],
]:
    by_basename: dict[str, list[SourceVariant]] = defaultdict(list)
    for variant in variants:
        by_basename[PurePosixPath(_normalize_path(variant.path)).name].append(
            variant
        )
    matched: list[tuple[_SourceDate, SourceVariant]] = []
    unmatched: list[str] = []
    ambiguous: list[dict[str, Any]] = []
    for source in source_dates.values():
        normalized = _normalize_path(source.source_path)
        if "/" in normalized:
            candidates = [
                variant
                for variant in variants
                if _normalize_path(variant.path).endswith(normalized)
            ]
        else:
            candidates = by_basename.get(PurePosixPath(normalized).name, [])
        if len(candidates) == 1:
            matched.append((source, candidates[0]))
        elif not candidates:
            unmatched.append(source.source_path)
        else:
            ambiguous.append({
                "source_path": source.source_path,
                "candidate_paths": [row.path for row in candidates],
            })
    return matched, unmatched, ambiguous


def _parse_date(value: str) -> date | None:
    try:
        parsed = date.fromisoformat(value)
    except ValueError:
        return None
    return parsed if 2000 <= parsed.year <= 2100 else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_path(value: str) -> str:
    return value.replace("\\", "/").strip("/").casefold()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _evidence_dict(
    current: DocumentMetadataVersion | None,
) -> dict[str, Any]:
    if current is None or not current.evidence_json:
        return {}
    try:
        value = json.loads(current.evidence_json)
    except (json.JSONDecodeError, TypeError):
        return {}
    return dict(value) if isinstance(value, dict) else {}
