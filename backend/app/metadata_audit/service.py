"""Extract document metadata only when the source file provides evidence."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.catalog.models import (
    DocumentMetadataCandidate,
    DocumentMetadataScan,
    DocumentMetadataVersion,
)
from app.catalog.service import current_document_metadata
from app.documents.models import SourceVariant
from app.quotes.models import RawQuoteItem
from app.standard_database.models import (
    QuoteDocumentPurpose,
    QuoteDocumentRole,
)


AUDIT_RULE_VERSION = "document-metadata-v4"
AUTOMATED_METADATA_ACTORS = {
    "metadata-audit-v1",
    "metadata-audit-v2",
    "team-standard-date-backfill-v1",
    "metadata-audit-v3",
    "metadata-audit-v4",
}
SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".pdf"}
PLACEHOLDER_VALUES = {
    "",
    "-",
    "미상",
    "미확인",
    "협력사미상",
    "aone협력사",
    "바츠추출",
}
PLACEHOLDER_DATES = {"2025-01-01"}
LABELS = {
    "supplier_name": (
        "공급사",
        "업체명",
        "회사명",
        "상호",
        "견적업체",
        "제출업체",
        "supplier",
        "vendor",
    ),
    "quote_date": (
        "견적일",
        "견적일자",
        "작성일",
        "작성일자",
        "발행일",
        "발행일자",
        "제출일",
        "제출일자",
        "提出日",
        "西紀",
        "서기",
        "quote date",
        "quotation date",
        "date of quote",
        "issue date",
        "issued date",
        "quote update",
    ),
    "project_name": (
        "공사명",
        "프로젝트명",
        "project name",
        "project",
    ),
}
DATE_PATTERNS = (
    "%Y-%m-%d",
    "%Y.%m.%d",
    "%Y/%m/%d",
    "%Y년 %m월 %d일",
)
DATE_REGEXES = (
    re.compile(
        r"(?<!\d)(20\d{2})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*"
        r"(\d{1,2})(?!\d)"
    ),
    re.compile(
        r"(?<!\d)(20\d{2})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일?"
    ),
    re.compile(
        r"(?<![A-Za-z0-9])(20\d{2})(\d{2})(\d{2})(?![A-Za-z0-9])"
    ),
)
ENGLISH_DATE_REGEXES = (
    (
        re.compile(
            r"(?i)\b(?:January|February|March|April|May|June|July|August|"
            r"September|October|November|December)\s+\d{1,2},\s+20\d{2}\b"
        ),
        "%B %d, %Y",
    ),
    (
        re.compile(
            r"(?i)\b\d{1,2}-(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|"
            r"Nov|Dec)-(?:20\d{2}|\d{2})\b"
        ),
        None,
    ),
)
QUOTE_DATE_LABEL = re.compile(
    r"(?i)(견\s*적\s*(?:일|일자|날짜)|작\s*성\s*(?:일|일자)|"
    r"발\s*행\s*(?:일|일자)|제\s*출\s*(?:일|일자)|"
    r"提\s*出\s*日|서\s*기|西\s*紀|quote\s*update|"
    r"quotation\s*date|quote\s*date|date\s*of\s*quote|"
    r"issue(?:d)?\s*date|(?:^|\s)date\s*[:：])"
)
NON_QUOTE_DATE_LABEL = re.compile(
    r"(?i)(납기|납품|유효|견적\s*유효|공사\s*기간|작업\s*기간|"
    r"준공|계약|발주|delivery|valid(?:ity)?|due\s*date|payment)"
)
QUOTE_TITLE = re.compile(
    r"(?i)(견\s*적\s*서|見\s*積\s*書|quotation|\bquote\b)"
)
COMPANY_NAME_MARKER = re.compile(
    r"(?i)(?:주식회사|유한회사|㈜|\(\s*주\s*\)|（\s*주\s*）|"
    r"\bco\.?\s*,?\s*ltd\.?\b|\bcorporation\b|\bcorp\.?\b)"
)
NON_SUPPLIER_HEADER_TEXT = re.compile(
    r"(?i)(?:주소|대표|담당|전화|연락처|fax|견적일|견적번호|현대|기아|貴下)"
)


@dataclass(frozen=True)
class ExtractedCandidate:
    field_name: str
    value_text: str
    source_kind: str
    confidence: int
    sheet: str | None = None
    page: int | None = None
    cells: str | None = None
    context_excerpt: str | None = None


@dataclass(frozen=True)
class FileAuditResult:
    relative_path: str
    extension: str
    acquisition_channel: str | None
    open_status: str
    content_status: str
    review_status: str
    supplier_name: str | None
    quote_date: str | None
    project_name: str | None
    source_locations: str
    confidence: int | None
    diagnostic: str | None


@dataclass(frozen=True)
class MetadataAuditReport:
    total_files: int
    opened_files: int
    failed_files: int
    unsupported_files: int
    auto_confirmed_files: int
    review_required_files: int
    supplier_confirmed_files: int
    date_confirmed_files: int
    report_file: str
    rule_version: str = AUDIT_RULE_VERSION

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def audit_quote_metadata(
    session: Session,
    *,
    quote_root: Path,
    report_path: Path,
    third_training_relative: Path = Path("3차 학습"),
    all_historical: bool = False,
) -> MetadataAuditReport:
    """Audit source files and append only source-backed metadata evidence."""

    quote_root = quote_root.resolve(strict=True)
    variants = {
        _path_key(row.path): row
        for row in session.scalars(select(SourceVariant).order_by(SourceVariant.id))
    }
    if all_historical:
        audit_root = quote_root
        target_variants = _preferred_historical_variants(session)
        targets = []
        for variant in target_variants:
            path = (quote_root / variant.path).resolve(strict=False)
            try:
                path.relative_to(quote_root)
            except ValueError:
                continue
            if path.is_file():
                targets.append((path, variant.path, variant.path, variant))
    else:
        audit_root = (quote_root / third_training_relative).resolve(strict=False)
        audit_root.relative_to(quote_root)
        targets = (
            [
                (
                    path,
                    path.relative_to(quote_root).as_posix(),
                    path.relative_to(audit_root).as_posix(),
                    variants.get(
                        _path_key(path.relative_to(quote_root).as_posix())
                    ),
                )
                for path in sorted(
                    (row for row in audit_root.rglob("*") if row.is_file()),
                    key=lambda row: row.as_posix().casefold(),
                )
            ]
            if audit_root.is_dir()
            else []
        )
    if not targets:
        _write_report(report_path, [])
        return MetadataAuditReport(
            total_files=0,
            opened_files=0,
            failed_files=0,
            unsupported_files=0,
            auto_confirmed_files=0,
            review_required_files=0,
            supplier_confirmed_files=0,
            date_confirmed_files=0,
            report_file=str(report_path),
        )
    existing_scans = {
        (row.source_path, row.input_fingerprint, row.rule_version): row
        for row in session.scalars(select(DocumentMetadataScan))
    }
    results: list[FileAuditResult] = []
    for path, relative_to_quote, relative_to_audit, variant in targets:
        fingerprint = _sha256(path)
        existing = existing_scans.get(
            (relative_to_quote, fingerprint, AUDIT_RULE_VERSION)
        )
        if existing is not None:
            restored = _stored_result(existing)
            if restored is not None:
                results.append(restored)
                continue
        result, candidates = _audit_file(path, relative_to_audit)
        _store_scan(
            session,
            variant,
            source_path=relative_to_quote,
            input_fingerprint=fingerprint,
            result=result,
            candidates=candidates,
        )
        results.append(result)
    session.flush()
    _write_report(report_path, results)
    return MetadataAuditReport(
        total_files=len(results),
        opened_files=sum(row.open_status == "OPENED" for row in results),
        failed_files=sum(row.open_status == "FAILED" for row in results),
        unsupported_files=sum(row.open_status == "UNSUPPORTED" for row in results),
        auto_confirmed_files=sum(row.review_status == "AUTO_CONFIRMED" for row in results),
        review_required_files=sum(row.review_status == "REVIEW_REQUIRED" for row in results),
        supplier_confirmed_files=sum(row.supplier_name is not None for row in results),
        date_confirmed_files=sum(row.quote_date is not None for row in results),
        report_file=str(report_path),
    )


def _preferred_historical_variants(session: Session) -> list[SourceVariant]:
    """Return one evidence variant per historical document.

    A variant that actually owns parsed rows is preferred, followed by the
    explicitly selected parse variant. This prevents protected/unlocked copies
    of the same document from creating duplicate metadata versions.
    """

    latest_roles = (
        select(
            QuoteDocumentRole.document_id.label("document_id"),
            func.max(QuoteDocumentRole.id).label("role_id"),
        )
        .group_by(QuoteDocumentRole.document_id)
        .subquery()
    )
    historical_ids = set(
        session.scalars(
            select(QuoteDocumentRole.document_id)
            .join(latest_roles, latest_roles.c.role_id == QuoteDocumentRole.id)
            .where(
                QuoteDocumentRole.purpose
                == QuoteDocumentPurpose.HISTORICAL_REFERENCE
            )
        )
    )
    parsed_variant_ids = set(
        session.scalars(select(RawQuoteItem.source_variant_id).distinct())
    )
    grouped: dict[int, list[SourceVariant]] = {}
    for variant in session.scalars(
        select(SourceVariant)
        .where(SourceVariant.document_id.in_(historical_ids))
        .order_by(SourceVariant.id)
    ):
        grouped.setdefault(variant.document_id, []).append(variant)
    return [
        max(
            rows,
            key=lambda row: (
                row.id in parsed_variant_ids,
                row.selected_for_parsing_at_ingest,
                row.id,
            ),
        )
        for _, rows in sorted(grouped.items())
    ]


def _audit_file(
    path: Path,
    relative_path: str,
) -> tuple[FileAuditResult, list[ExtractedCandidate]]:
    extension = path.suffix.lower()
    channel = _acquisition_channel(relative_path)
    if extension not in SUPPORTED_EXTENSIONS:
        return (
            FileAuditResult(
                relative_path=relative_path,
                extension=extension,
                acquisition_channel=channel,
                open_status="UNSUPPORTED",
                content_status="MANUAL_REVIEW",
                review_status="REVIEW_REQUIRED",
                supplier_name=None,
                quote_date=None,
                project_name=None,
                source_locations="",
                confidence=None,
                diagnostic="unsupported evidence format",
            ),
            [],
        )
    try:
        if extension == ".xlsx":
            candidates, has_text = _xlsx_candidates(path)
        elif extension == ".xls":
            candidates, has_text = _xls_candidates(path)
        else:
            candidates, has_text = _pdf_candidates(path)
    except Exception as exc:  # file-level audit must continue
        return (
            FileAuditResult(
                relative_path=relative_path,
                extension=extension,
                acquisition_channel=channel,
                open_status="FAILED",
                content_status="UNREADABLE",
                review_status="REVIEW_REQUIRED",
                supplier_name=None,
                quote_date=None,
                project_name=None,
                source_locations="",
                confidence=None,
                diagnostic=type(exc).__name__,
            ),
            [],
        )
    candidates = _deduplicate_candidates(candidates)
    selected, ambiguous_fields = _select_candidates(candidates)
    locations = "; ".join(
        _candidate_location(row) for row in selected.values()
    )
    review_status = (
        "AUTO_CONFIRMED"
        if {"supplier_name", "quote_date"}.issubset(selected)
        and not ({"supplier_name", "quote_date"} & ambiguous_fields)
        else "REVIEW_REQUIRED"
    )
    return (
        FileAuditResult(
            relative_path=relative_path,
            extension=extension,
            acquisition_channel=channel,
            open_status="OPENED",
            content_status="TEXT" if has_text else "OCR_REQUIRED",
            review_status=review_status,
            supplier_name=(
                selected.get("supplier_name").value_text
                if "supplier_name" in selected else None
            ),
            quote_date=(
                selected.get("quote_date").value_text
                if "quote_date" in selected else None
            ),
            project_name=(
                selected.get("project_name").value_text
                if "project_name" in selected else None
            ),
            source_locations=locations,
            confidence=(
                min(row.confidence for row in selected.values())
                if selected else None
            ),
            diagnostic=(
                "conflicting source values: "
                + ", ".join(sorted(ambiguous_fields))
                if ambiguous_fields else None
            ),
        ),
        candidates,
    )


def _deduplicate_candidates(
    candidates: Iterable[ExtractedCandidate],
) -> list[ExtractedCandidate]:
    result: dict[tuple[str, str, str | None, int | None], ExtractedCandidate] = {}
    for candidate in candidates:
        key = (
            candidate.field_name,
            candidate.value_text.casefold(),
            candidate.sheet,
            candidate.page,
        )
        current = result.get(key)
        if current is None or candidate.confidence > current.confidence:
            result[key] = candidate
    return list(result.values())


def _xlsx_candidates(path: Path) -> tuple[list[ExtractedCandidate], bool]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        candidates: list[ExtractedCandidate] = []
        has_text = False
        for sheet in workbook.worksheets[:20]:
            rows = [
                list(row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 120),
                    min_col=1,
                    max_col=min(sheet.max_column, 32),
                    values_only=True,
                )
            ]
            has_text = has_text or any(
                str(value).strip()
                for row in rows for value in row if value is not None
            )
            candidates.extend(_grid_candidates(rows, sheet.title))
        return candidates, has_text
    finally:
        workbook.close()


def _xls_candidates(path: Path) -> tuple[list[ExtractedCandidate], bool]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    candidates: list[ExtractedCandidate] = []
    has_text = False
    try:
        for sheet in workbook.sheets()[:20]:
            rows = [
                [
                    (
                        xlrd.xldate_as_datetime(
                            sheet.cell_value(row, col),
                            workbook.datemode,
                        )
                        if sheet.cell_type(row, col) == xlrd.XL_CELL_DATE
                        else sheet.cell_value(row, col)
                    )
                    for col in range(min(sheet.ncols, 32))
                ]
                for row in range(min(sheet.nrows, 120))
            ]
            has_text = has_text or any(
                str(value).strip()
                for row in rows for value in row if value is not None
            )
            candidates.extend(_grid_candidates(rows, sheet.name))
        return candidates, has_text
    finally:
        workbook.release_resources()


def _pdf_candidates(path: Path) -> tuple[list[ExtractedCandidate], bool]:
    reader = PdfReader(str(path))
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise ValueError("encrypted pdf")
    candidates: list[ExtractedCandidate] = []
    has_text = False
    for page_number, page in enumerate(reader.pages[:5], start=1):
        text = page.extract_text() or ""
        if text.strip():
            has_text = True
        for line in text.splitlines():
            candidates.extend(_line_candidates(line, page=page_number))
        candidates.extend(
            _date_text_candidates(
                text,
                page=page_number,
                quote_header=(page_number == 1 and bool(QUOTE_TITLE.search(text))),
            )
        )
    if not has_text:
        try:
            from app.ingestion.readers import (
                OcrReviewRequiredError,
                OcrUnavailableError,
                UnsafeQuoteFileError,
                ocr_pdf_text_pages,
            )

            ocr_pages = ocr_pdf_text_pages(path, page_limit=2)
        except (OcrReviewRequiredError, OcrUnavailableError, UnsafeQuoteFileError):
            ocr_pages = ()
        for page_number, text in enumerate(ocr_pages, start=1):
            ocr_candidates = _date_text_candidates(
                text,
                page=page_number,
                quote_header=(
                    page_number == 1 and bool(QUOTE_TITLE.search(text))
                ),
                ocr_source=True,
            )
            candidates.extend(ocr_candidates)
    return candidates, has_text


def _grid_candidates(
    rows: list[list[object]],
    sheet: str,
) -> list[ExtractedCandidate]:
    candidates: list[ExtractedCandidate] = []
    for row_index, row in enumerate(rows, start=1):
        row_text = " | ".join(
            str(value) for value in row if value is not None
        )
        candidates.extend(
            _date_text_candidates(
                row_text,
                sheet=sheet,
                cells=f"A{row_index}:{get_column_letter(len(row))}{row_index}",
            )
        )
        for col_index, raw_value in enumerate(row, start=1):
            if raw_value is None:
                continue
            text = str(raw_value).strip()
            if row_index <= 15:
                supplier = _quote_header_supplier(text)
                if supplier is not None:
                    candidates.append(
                        ExtractedCandidate(
                            field_name="supplier_name",
                            value_text=supplier,
                            source_kind="QUOTE_HEADER_COMPANY_NAME",
                            confidence=94,
                            sheet=sheet,
                            cells=(
                                f"{get_column_letter(col_index)}{row_index}"
                            ),
                        )
                    )
            if isinstance(raw_value, (date, datetime)):
                typed_date = (
                    raw_value.date()
                    if isinstance(raw_value, datetime)
                    else raw_value
                )
                if 2000 <= typed_date.year <= date.today().year:
                    left = " ".join(
                        str(value)
                        for value in row[max(0, col_index - 4):col_index - 1]
                        if value is not None
                    )
                    if QUOTE_DATE_LABEL.search(left) and not NON_QUOTE_DATE_LABEL.search(left):
                        candidates.append(
                            ExtractedCandidate(
                                field_name="quote_date",
                                value_text=typed_date.isoformat(),
                                source_kind="EXPLICIT_QUOTE_DATE_CELL",
                                confidence=99,
                                sheet=sheet,
                                cells=f"{get_column_letter(col_index)}{row_index}",
                                context_excerpt=f"{left} | {typed_date.isoformat()}",
                            )
                        )
            candidates.extend(
                _line_candidates(
                    text,
                    sheet=sheet,
                    cells=f"{get_column_letter(col_index)}{row_index}",
                )
            )
            field_name = _label_field(text)
            if field_name is None:
                continue
            adjacent = _next_value(row, col_index)
            value = _clean_candidate(field_name, adjacent)
            if value is None:
                continue
            end_col = min(len(row), col_index + 2)
            candidates.append(
                ExtractedCandidate(
                    field_name=field_name,
                    value_text=value,
                    source_kind="LABELED_CELL",
                    confidence=95,
                    sheet=sheet,
                    cells=(
                        f"{get_column_letter(col_index)}{row_index}:"
                        f"{get_column_letter(end_col)}{row_index}"
                    ),
                )
            )
    return candidates


def _quote_header_supplier(value: str) -> str | None:
    """Accept an unlabeled legal company name in a quote cover header."""

    cleaned = " ".join(value.split()).strip(" :：")
    if not 2 <= len(cleaned) <= 100:
        return None
    if NON_SUPPLIER_HEADER_TEXT.search(cleaned):
        return None
    if COMPANY_NAME_MARKER.search(cleaned) is None:
        return None
    return _clean_candidate("supplier_name", cleaned)


def _date_text_candidates(
    text: str,
    *,
    sheet: str | None = None,
    page: int | None = None,
    cells: str | None = None,
    quote_header: bool = False,
    ocr_source: bool = False,
) -> list[ExtractedCandidate]:
    """Extract source-confirmed quote dates without guessing from filenames."""

    compact = " ".join(text.split())
    mentions: list[tuple[date, bool, bool, str]] = []
    seen: set[tuple[date, int]] = set()
    for parsed, match_start, match_end in _date_occurrences(compact):
        if not (2000 <= parsed.year <= date.today().year):
            continue
        if parsed.isoformat() in PLACEHOLDER_DATES:
            continue
        key = (parsed, match_start)
        if key in seen:
            continue
        seen.add(key)
        before = compact[max(0, match_start - 70):match_start]
        context = compact[
            max(0, match_start - 80):min(len(compact), match_end + 80)
        ]
        labels = list(QUOTE_DATE_LABEL.finditer(before))
        explicit = False
        if labels:
            gap = before[labels[-1].end():]
            explicit = (
                len(gap) <= 24
                and re.fullmatch(r"[\s:：()\[\].\-/|]*", gap) is not None
            )
        negative = bool(NON_QUOTE_DATE_LABEL.search(context)) and not explicit
        mentions.append((parsed, explicit, negative, context))

    result: list[ExtractedCandidate] = []
    for parsed, explicit, negative, context in mentions:
        if not explicit or negative:
            continue
        result.append(
            ExtractedCandidate(
                field_name="quote_date",
                value_text=parsed.isoformat(),
                source_kind=(
                    "OCR_EXPLICIT_QUOTE_DATE_TEXT"
                    if ocr_source else "EXPLICIT_QUOTE_DATE_TEXT"
                ),
                confidence=90 if ocr_source else 99,
                sheet=sheet,
                page=page,
                cells=cells,
                context_excerpt=context,
            )
        )
    if quote_header:
        unlabeled_dates = {
            parsed for parsed, explicit, negative, _ in mentions
            if not explicit and not negative
        }
        if len(unlabeled_dates) == 1:
            header_date = next(iter(unlabeled_dates))
            header_context = next(
                context for parsed, explicit, negative, context in mentions
                if parsed == header_date and not explicit and not negative
            )
            result.append(
                ExtractedCandidate(
                    field_name="quote_date",
                    value_text=header_date.isoformat(),
                    source_kind=(
                        "OCR_QUOTE_HEADER_DATE"
                        if ocr_source else "QUOTE_HEADER_DATE"
                    ),
                    confidence=90,
                    sheet=sheet,
                    page=page,
                    cells=cells,
                    context_excerpt=header_context,
                )
            )
    return result


def _date_occurrences(text: str) -> list[tuple[date, int, int]]:
    result: list[tuple[date, int, int]] = []
    for pattern in DATE_REGEXES:
        for match in pattern.finditer(text):
            try:
                parsed = date(*(int(part) for part in match.groups()))
            except ValueError:
                continue
            result.append((parsed, match.start(), match.end()))
    for pattern, date_format in ENGLISH_DATE_REGEXES:
        for match in pattern.finditer(text):
            formats = (
                (date_format,)
                if date_format is not None
                else ("%d-%b-%Y", "%d-%b-%y")
            )
            parsed = None
            for candidate_format in formats:
                try:
                    parsed = datetime.strptime(
                        match.group(0), candidate_format
                    ).date()
                    break
                except ValueError:
                    continue
            if parsed is not None:
                result.append((parsed, match.start(), match.end()))
    return result


def _line_candidates(
    line: str,
    *,
    sheet: str | None = None,
    page: int | None = None,
    cells: str | None = None,
) -> list[ExtractedCandidate]:
    compact = " ".join(line.split())
    result: list[ExtractedCandidate] = []
    for field_name, aliases in LABELS.items():
        for alias in aliases:
            match = re.search(
                rf"(?i)(?:^|\s){re.escape(alias)}\s*[:：]\s*(.+)$",
                compact,
            )
            if match is None:
                continue
            value = _clean_candidate(field_name, match.group(1))
            if value is not None:
                result.append(
                    ExtractedCandidate(
                        field_name=field_name,
                        value_text=value,
                        source_kind="LABELED_TEXT",
                        confidence=95,
                        sheet=sheet,
                        page=page,
                        cells=cells,
                    )
                )
            break
    return result


def _label_field(value: str) -> str | None:
    normalized = re.sub(r"[\s:：()\[\]]+", "", value).casefold()
    for field_name, aliases in LABELS.items():
        if normalized in {
            re.sub(r"[\s:：()\[\]]+", "", alias).casefold()
            for alias in aliases
        }:
            return field_name
    return None


def _next_value(row: list[object], one_based_col: int) -> str:
    for value in row[one_based_col: one_based_col + 2]:
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _clean_candidate(field_name: str, value: str) -> str | None:
    cleaned = " ".join(str(value).split()).strip(" :：")
    if not cleaned or len(cleaned) > 300:
        return None
    if cleaned.casefold() in PLACEHOLDER_VALUES:
        return None
    if field_name == "supplier_name":
        if cleaned.casefold() in {"aone", "바츠"}:
            return None
        if not re.search(r"[A-Za-z가-힣]", cleaned):
            return None
    if field_name == "quote_date":
        parsed = _parse_date(cleaned)
        if parsed is None or parsed.isoformat() in PLACEHOLDER_DATES:
            return None
        return parsed.isoformat()
    return cleaned


def _parse_date(value: str) -> date | None:
    normalized = value.strip()
    for pattern in DATE_PATTERNS:
        try:
            return datetime.strptime(normalized, pattern).date()
        except ValueError:
            continue
    for pattern in DATE_REGEXES:
        match = pattern.search(normalized)
        if match is None:
            continue
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            continue
    return None


def _select_candidates(
    candidates: Iterable[ExtractedCandidate],
) -> tuple[dict[str, ExtractedCandidate], set[str]]:
    by_field: dict[str, list[ExtractedCandidate]] = {}
    for candidate in candidates:
        by_field.setdefault(candidate.field_name, []).append(candidate)
    selected: dict[str, ExtractedCandidate] = {}
    ambiguous_fields: set[str] = set()
    for field_name, rows in by_field.items():
        best_confidence = max(row.confidence for row in rows)
        preferred = [row for row in rows if row.confidence == best_confidence]
        if field_name == "quote_date" and any(
            row.page is not None for row in preferred
        ):
            first_page = min(
                row.page for row in preferred if row.page is not None
            )
            preferred = [row for row in preferred if row.page == first_page]
        values = {row.value_text.casefold() for row in preferred}
        if len(values) != 1:
            ambiguous_fields.add(field_name)
            continue
        selected[field_name] = sorted(
            preferred,
            key=lambda row: (-row.confidence, row.page or 0, row.cells or ""),
        )[0]
    return selected, ambiguous_fields


def _store_scan(
    session: Session,
    variant: SourceVariant | None,
    *,
    source_path: str,
    input_fingerprint: str,
    result: FileAuditResult,
    candidates: list[ExtractedCandidate],
) -> DocumentMetadataScan:
    existing = session.scalar(
        select(DocumentMetadataScan).where(
            DocumentMetadataScan.source_path == source_path,
            DocumentMetadataScan.input_fingerprint == input_fingerprint,
            DocumentMetadataScan.rule_version == AUDIT_RULE_VERSION,
        )
    )
    if existing is not None:
        return existing
    diagnostics = {
        "diagnostic": result.diagnostic,
        "relative_path": result.relative_path,
        "result": asdict(result),
    }
    scan = DocumentMetadataScan(
        source_variant_id=None if variant is None else variant.id,
        source_path=source_path,
        rule_version=AUDIT_RULE_VERSION,
        input_fingerprint=input_fingerprint,
        open_status=result.open_status,
        content_status=result.content_status,
        review_status=result.review_status,
        acquisition_channel=result.acquisition_channel,
        diagnostics_json=_json(diagnostics),
    )
    session.add(scan)
    session.flush()
    selected, ambiguous_fields = _select_candidates(candidates)
    candidate_rows: list[DocumentMetadataCandidate] = []
    for candidate in candidates:
        accepted = (
            candidate.field_name not in ambiguous_fields
            and selected.get(candidate.field_name) == candidate
            and candidate.confidence >= 90
        )
        evidence = {
            "sheet": candidate.sheet,
            "page": candidate.page,
            "cells": candidate.cells,
            "source_kind": candidate.source_kind,
            "context_excerpt": candidate.context_excerpt,
        }
        fingerprint = hashlib.sha256(
            _json(
                {
                    "field": candidate.field_name,
                    "value": candidate.value_text,
                    **evidence,
                }
            ).encode("utf-8")
        ).hexdigest()
        row = DocumentMetadataCandidate(
            scan_id=scan.id,
            field_name=candidate.field_name,
            value_text=candidate.value_text,
            source_kind=candidate.source_kind,
            confidence=candidate.confidence,
            status="AUTO_ACCEPTED" if accepted else "REVIEW_REQUIRED",
            source_sheet=candidate.sheet,
            source_page=candidate.page,
            source_cells=candidate.cells,
            evidence_json=_json(evidence),
            candidate_fingerprint=fingerprint,
        )
        session.add(row)
        candidate_rows.append(row)
    session.flush()
    # A standalone ingest can own parsed rows even though it was not marked as
    # the preferred member of a protected/unlocked source group.  In either
    # case the metadata belongs to the exact variant that supplied the rows.
    if variant is not None and (
        variant.selected_for_parsing_at_ingest or bool(variant.raw_items)
    ):
        _append_auto_metadata(session, variant, candidate_rows)
    return scan


def _append_auto_metadata(
    session: Session,
    variant: SourceVariant,
    candidates: list[DocumentMetadataCandidate],
) -> None:
    accepted = {
        row.field_name: row
        for row in candidates
        if row.status == "AUTO_ACCEPTED"
    }
    if not accepted:
        return
    current = current_document_metadata(session, variant.document_id)
    if (
        current is not None
        and current.decided_by not in AUTOMATED_METADATA_ACTORS
    ):
        return
    supplier = accepted.get("supplier_name")
    quote_date = accepted.get("quote_date")
    project = accepted.get("project_name")
    values = {
        "supplier_name": (
            supplier.value_text
            if supplier
            else (current.supplier_name if current is not None else None)
        ),
        "quote_date": (
            date.fromisoformat(quote_date.value_text)
            if quote_date
            else (current.quote_date if current is not None else None)
        ),
        "project_name": (
            project.value_text
            if project
            else (current.project_name if current is not None else None)
        ),
    }
    try:
        prior_evidence = (
            json.loads(current.evidence_json)
            if current is not None and current.evidence_json
            else {}
        )
    except (json.JSONDecodeError, TypeError):
        prior_evidence = {}
    if current is not None and all(
        getattr(current, field) == value for field, value in values.items()
    ) and all(
        isinstance(prior_evidence, dict)
        and isinstance(prior_evidence.get(field), dict)
        and prior_evidence[field].get("quality") == "SOURCE_CONFIRMED"
        for field in accepted
    ):
        return
    evidence = dict(prior_evidence) if isinstance(prior_evidence, dict) else {}
    evidence.update({
        field: {
            "candidate_id": row.id,
            "source_kind": row.source_kind,
            "confidence": row.confidence,
            "sheet": row.source_sheet,
            "page": row.source_page,
            "cells": row.source_cells,
            "quality": "SOURCE_CONFIRMED",
            "use_for_index": "EXACT_DATE" if field == "quote_date" else None,
        }
        for field, row in accepted.items()
    })
    session.add(
        DocumentMetadataVersion(
            source_document_id=variant.document_id,
            version_number=1 if current is None else current.version_number + 1,
            **values,
            decided_by="metadata-audit-v4",
            reason_detail="원본 견적서 본문·머리말에서 자동 확인",
            evidence_json=_json(evidence),
        )
    )


def _write_report(path: Path, rows: list[FileAuditResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(asdict(rows[0]).keys()) if rows else ["relative_path"])
        writer.writeheader()
        for row in rows:
            writer.writerow(asdict(row))
    temporary.replace(path)


def _candidate_location(candidate: ExtractedCandidate) -> str:
    if candidate.sheet:
        return f"{candidate.field_name}:{candidate.sheet}!{candidate.cells or '?'}"
    if candidate.page:
        return f"{candidate.field_name}:PDF {candidate.page}쪽"
    return candidate.field_name


def _acquisition_channel(path: str) -> str | None:
    normalized = path.casefold()
    if "aone" in normalized:
        return "AONE"
    if "바츠" in normalized:
        return "BATS"
    return None


def _path_key(path: str) -> str:
    return path.replace("\\", "/").casefold()


def _json(payload: object) -> str:
    return json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _stored_result(scan: DocumentMetadataScan) -> FileAuditResult | None:
    try:
        payload = json.loads(scan.diagnostics_json)
        result = payload.get("result")
        if not isinstance(result, dict):
            return None
        return FileAuditResult(**result)
    except (json.JSONDecodeError, TypeError, KeyError, ValueError):
        return None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
