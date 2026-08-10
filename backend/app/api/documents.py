"""Source-document inventory and configured-folder ingestion."""

from __future__ import annotations

import hashlib
import mimetypes
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
import xlrd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.cleansing.models import CleanDecision, CleanStatus
from app.core.config import settings
from app.db.session import get_session
from app.documents.models import SourceDocument, SourceVariant
from app.ingestion.corpus import ingest_corpus
from app.ingestion.service import preferred_variant_for
from app.quotes.models import RawQuoteItem


router = APIRouter()


class VariantResponse(BaseModel):
    id: int
    path: str
    sha256: str
    extension: str
    security_state: str
    selected_for_parsing_at_ingest: bool
    registered_at: datetime
    raw_item_count: int


class DocumentCountsResponse(BaseModel):
    raw_items: int
    INCLUDED: int
    EXCLUDED: int
    REVIEW_REQUIRED: int
    UNDECIDED: int


class DocumentResponse(BaseModel):
    id: int
    logical_name: str
    created_at: datetime
    variants: list[VariantResponse]
    preferred_variant: VariantResponse
    counts: DocumentCountsResponse


class DocumentListResponse(BaseModel):
    items: list[DocumentResponse]
    total: int
    limit: int
    offset: int


class VariantEvidenceResponse(BaseModel):
    path: str
    sha256: str | None = None
    error_code: str | None = None


class ScanFailureResponse(BaseModel):
    logical_name: str
    error_code: str
    detail: str
    preferred_path: str | None = None
    preferred_sha256: str | None = None
    variants: list[VariantEvidenceResponse] | None = None


class ScanResponse(BaseModel):
    files_found: int
    documents_found: int
    documents_succeeded: int
    documents_failed: int
    documents_review_required: int
    variants_created: int
    raw_items_created: int
    decisions_created: int
    failures: list[ScanFailureResponse]
    review_required: list[ScanFailureResponse]


class PreviewCellResponse(BaseModel):
    coordinate: str
    value: str | None
    highlighted: bool


class PreviewRowResponse(BaseModel):
    row_number: int
    cells: list[PreviewCellResponse]


class VariantPreviewResponse(BaseModel):
    kind: str
    file_url: str
    file_name: str
    sheet: str | None = None
    page: int | None = None
    target_cells: str | None = None
    rows: list[PreviewRowResponse] = Field(default_factory=list)


@router.get("", response_model=DocumentListResponse)
def list_documents(
    session: Session = Depends(get_session),
    *,
    limit: int = Query(50, ge=1, le=100),
    offset: int = Query(0, ge=0),
) -> dict[str, object]:
    total = session.scalar(select(func.count(SourceDocument.id))) or 0
    documents = session.scalars(
        select(SourceDocument)
        .options(
            selectinload(SourceDocument.variants).selectinload(
                SourceVariant.raw_items
            )
        )
        .order_by(SourceDocument.logical_name, SourceDocument.id)
        .offset(offset)
        .limit(limit)
    ).all()
    document_ids = [document.id for document in documents]
    current_by_item = _current_decisions(session, document_ids)
    return {
        "items": [
            _document_item(document, current_by_item)
            for document in documents
        ],
        "total": total,
        "limit": limit,
        "offset": offset,
    }


@router.post(
    "/scan",
    response_model=ScanResponse,
    response_model_exclude_none=True,
)
def scan_documents(
    session: Session = Depends(get_session),
) -> dict[str, object]:
    report = ingest_corpus(session, settings.quote_path)
    return {
        "files_found": report.preflight.physical_files,
        "documents_found": report.preflight.logical_documents,
        "documents_succeeded": (
            report.documents_ingested + report.documents_unchanged
        ),
        "documents_failed": report.documents_failed,
        "documents_review_required": report.documents_review_required,
        "variants_created": report.variants_created,
        "raw_items_created": report.raw_items_created,
        "decisions_created": (
            report.base_decisions_created
            + report.outlier_decisions_created
        ),
        "failures": [failure.to_dict() for failure in report.failures],
        "review_required": [
            issue.to_dict() for issue in report.review_required
        ],
    }


@router.get("/variants/{variant_id}/file", response_class=FileResponse)
def get_variant_file(
    variant_id: int,
    session: Session = Depends(get_session),
) -> FileResponse:
    variant = session.get(SourceVariant, variant_id)
    if variant is None:
        raise HTTPException(status_code=404, detail="source file not found")
    path = _resolve_variant_file(variant)
    if path is None:
        raise HTTPException(status_code=404, detail="source file not found")
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    disposition = f"inline; filename*=UTF-8''{quote(path.name)}"
    return FileResponse(
        path,
        media_type=media_type,
        headers={"Content-Disposition": disposition},
    )


@router.get(
    "/variants/{variant_id}/preview",
    response_model=VariantPreviewResponse,
)
def get_variant_preview(
    variant_id: int,
    raw_item_id: int = Query(gt=0),
    session: Session = Depends(get_session),
) -> dict[str, object]:
    variant = session.get(SourceVariant, variant_id)
    raw = session.get(RawQuoteItem, raw_item_id)
    if variant is None or raw is None or raw.source_variant_id != variant.id:
        raise HTTPException(status_code=404, detail="source evidence not found")
    path = _resolve_variant_file(variant)
    if path is None:
        raise HTTPException(status_code=404, detail="source file not found")
    file_url = f"/api/documents/variants/{variant.id}/file"
    extension = path.suffix.lower()
    if extension == ".pdf":
        return {
            "kind": "PDF",
            "file_url": file_url,
            "file_name": path.name,
            "page": raw.source_page or 1,
            "target_cells": None,
            "rows": [],
        }
    if extension not in {".xlsx", ".xls"}:
        return {
            "kind": "FILE",
            "file_url": file_url,
            "file_name": path.name,
            "page": raw.source_page,
            "target_cells": raw.source_cells,
            "rows": [],
        }
    try:
        rows = (
            _xlsx_preview(path, raw)
            if extension == ".xlsx"
            else _xls_preview(path, raw)
        )
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"source preview unavailable: {type(exc).__name__}",
        ) from exc
    return {
        "kind": "SPREADSHEET",
        "file_url": file_url,
        "file_name": path.name,
        "sheet": raw.source_sheet,
        "page": None,
        "target_cells": raw.source_cells,
        "rows": rows,
    }


def _current_decisions(
    session: Session,
    document_ids: list[int],
) -> dict[int, CleanDecision]:
    if not document_ids:
        return {}
    latest_ids = (
        select(
            CleanDecision.raw_item_id,
            func.max(CleanDecision.id).label("decision_id"),
        )
        .join(
            RawQuoteItem,
            RawQuoteItem.id == CleanDecision.raw_item_id,
        )
        .join(
            SourceVariant,
            SourceVariant.id == RawQuoteItem.source_variant_id,
        )
        .where(SourceVariant.document_id.in_(document_ids))
        .group_by(CleanDecision.raw_item_id)
        .subquery()
    )
    return {
        decision.raw_item_id: decision
        for decision in session.scalars(
            select(CleanDecision).join(
                latest_ids,
                CleanDecision.id == latest_ids.c.decision_id,
            )
        )
    }


def _document_item(
    document: SourceDocument,
    current_by_item: dict[int, CleanDecision],
) -> dict[str, object]:
    variants = sorted(document.variants, key=lambda variant: variant.path)
    preferred = preferred_variant_for(document)
    raw_items = [
        raw_item
        for variant in variants
        for raw_item in variant.raw_items
    ]
    counts = {status.value: 0 for status in CleanStatus}
    undecided = 0
    for raw_item in raw_items:
        decision = current_by_item.get(raw_item.id)
        if decision is None:
            undecided += 1
        else:
            counts[decision.status.value] += 1
    return {
        "id": document.id,
        "logical_name": document.logical_name,
        "created_at": document.created_at.isoformat(),
        "variants": [_variant_item(variant) for variant in variants],
        "preferred_variant": _variant_item(preferred),
        "counts": {
            "raw_items": len(raw_items),
            **counts,
            "UNDECIDED": undecided,
        },
    }


def _variant_item(variant: SourceVariant) -> dict[str, object]:
    return {
        "id": variant.id,
        "path": variant.path,
        "sha256": variant.sha256,
        "extension": variant.extension,
        "security_state": variant.security_state,
        "selected_for_parsing_at_ingest": (
            variant.selected_for_parsing_at_ingest
        ),
        "registered_at": variant.registered_at.isoformat(),
        "raw_item_count": len(variant.raw_items),
    }


def _resolve_variant_file(variant: SourceVariant) -> Path | None:
    relative = Path(variant.path)
    if relative.is_absolute() or ".." in relative.parts:
        return None
    for configured_root in (settings.quote_path, settings.submission_path):
        root = configured_root.resolve(strict=False)
        candidate = (root / relative).resolve(strict=False)
        try:
            candidate.relative_to(root)
        except ValueError:
            continue
        if candidate.is_file() and _sha256(candidate) == variant.sha256:
            return candidate
    return None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _preview_bounds(raw: RawQuoteItem) -> tuple[int, int, int, int]:
    target_row = raw.source_row or 1
    first_col, last_col = 1, 12
    if raw.source_cells:
        matches = re.findall(r"([A-Za-z]+)(\d+)", raw.source_cells)
        if matches:
            indexes = [column_index_from_string(column) for column, _ in matches]
            first_col = max(1, min(indexes) - 1)
            last_col = min(30, max(indexes) + 1)
    return max(1, target_row - 3), target_row + 3, first_col, last_col


def _target_coordinates(raw: RawQuoteItem) -> set[str]:
    if not raw.source_cells:
        return set()
    matches = re.findall(r"([A-Za-z]+)(\d+)", raw.source_cells)
    if not matches:
        return set()
    first_col = column_index_from_string(matches[0][0])
    last_col = column_index_from_string(matches[-1][0])
    first_row = int(matches[0][1])
    last_row = int(matches[-1][1])
    return {
        f"{get_column_letter(column)}{row}"
        for row in range(first_row, last_row + 1)
        for column in range(first_col, last_col + 1)
    }


def _xlsx_preview(path: Path, raw: RawQuoteItem) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook[raw.source_sheet]
            if raw.source_sheet in workbook.sheetnames
            else workbook.worksheets[0]
        )
        first_row, last_row, first_col, last_col = _preview_bounds(raw)
        target = _target_coordinates(raw)
        return [
            {
                "row_number": row_index,
                "cells": [
                    {
                        "coordinate": f"{get_column_letter(column_index)}{row_index}",
                        "value": _preview_text(sheet.cell(row_index, column_index).value),
                        "highlighted": (
                            f"{get_column_letter(column_index)}{row_index}" in target
                            or (not target and row_index == raw.source_row)
                        ),
                    }
                    for column_index in range(first_col, last_col + 1)
                ],
            }
            for row_index in range(first_row, min(last_row, sheet.max_row) + 1)
        ]
    finally:
        workbook.close()


def _xls_preview(path: Path, raw: RawQuoteItem) -> list[dict[str, object]]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = (
            workbook.sheet_by_name(raw.source_sheet)
            if raw.source_sheet in workbook.sheet_names()
            else workbook.sheet_by_index(0)
        )
        first_row, last_row, first_col, last_col = _preview_bounds(raw)
        target = _target_coordinates(raw)
        rows: list[dict[str, object]] = []
        for row_index in range(first_row, min(last_row, sheet.nrows) + 1):
            cells = []
            for column_index in range(first_col, min(last_col, sheet.ncols) + 1):
                coordinate = f"{get_column_letter(column_index)}{row_index}"
                cells.append(
                    {
                        "coordinate": coordinate,
                        "value": _preview_text(
                            sheet.cell_value(row_index - 1, column_index - 1)
                        ),
                        "highlighted": (
                            coordinate in target
                            or (not target and row_index == raw.source_row)
                        ),
                    }
                )
            rows.append({"row_number": row_index, "cells": cells})
        return rows
    finally:
        workbook.release_resources()


def _preview_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
