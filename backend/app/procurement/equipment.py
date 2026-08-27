"""Persist an equipment-cover projection for one immutable analysis run."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analysis.models import QuoteAnalysisLineResult
from app.analysis.target_price import AnalysisRunResult
from app.core.config import settings
from app.procurement.models import (
    QuoteAnalysisEquipmentGroup,
    QuoteAnalysisEquipmentLine,
)


KRW = Decimal("1")
_DETAIL_SHEET = re.compile(r"(?:단위)?(?:장비|설비)\s*[-_ ]?(\d+)$", re.IGNORECASE)
_NUMBERED_SHEET = re.compile(r"#\s*(\d+)$")


@dataclass(frozen=True)
class EquipmentDefinition:
    key: str
    name: str
    sheet: str
    cover_quote_amount: Decimal | None
    source_kind: str


def create_equipment_projection(
    session: Session,
    result: AnalysisRunResult,
) -> tuple[QuoteAnalysisEquipmentGroup, ...]:
    existing = tuple(
        session.scalars(
            select(QuoteAnalysisEquipmentGroup)
            .where(QuoteAnalysisEquipmentGroup.analysis_run_id == result.run_id)
            .order_by(QuoteAnalysisEquipmentGroup.id)
        )
    )
    if existing:
        return existing

    definitions = _equipment_definitions(result)
    by_sheet = {definition.sheet: definition for definition in definitions}
    has_structured_cover = bool(definitions)
    source_lines = {
        line.raw_item_id: line
        for line in result.analysis.lines
        if not has_structured_cover or line.source.sheet in by_sheet
    }
    target_by_raw = {line.raw_item_id: line for line in result.target_lines}
    stored_by_raw = {
        line.raw_item_id: line
        for line in session.scalars(
            select(QuoteAnalysisLineResult).where(
                QuoteAnalysisLineResult.analysis_run_id == result.run_id
            )
        )
    }
    grouped: dict[str, list[int]] = defaultdict(list)
    definitions_by_key = {definition.key: definition for definition in definitions}
    for raw_id, line in source_lines.items():
        definition = by_sheet.get(line.source.sheet or "")
        if definition is None:
            sheet_name = (line.source.sheet or "").strip()
            fallback_key = f"sheet:{sheet_name or 'all-items'}"
            definitions_by_key.setdefault(
                fallback_key,
                EquipmentDefinition(
                    fallback_key,
                    sheet_name or "전체 설비",
                    sheet_name,
                    None,
                    "LINE_ITEM_FALLBACK",
                ),
            )
            grouped[fallback_key].append(raw_id)
        else:
            grouped[definition.key].append(raw_id)

    rows: list[QuoteAnalysisEquipmentGroup] = []
    for key, raw_ids in grouped.items():
        definition = definitions_by_key[key]
        detail_quote = sum(
            (
                max(source_lines[raw_id].quote_amount or Decimal("0"), Decimal("0"))
                for raw_id in raw_ids
            ),
            Decimal("0"),
        )
        if detail_quote <= 0:
            # Preformatted workbooks frequently ship unused numbered sheets.
            # A cover amount alone does not create an equipment group because
            # the detail sheet is the operational source of the calculation.
            continue
        raw_negotiation = sum(
            (_negotiation_amount(source_lines[raw_id].quote_amount, target_by_raw[raw_id].target_amount)
             for raw_id in raw_ids),
            Decimal("0"),
        ).quantize(KRW, rounding=ROUND_HALF_UP)
        # One detail sheet represents one equipment.  The equipment totals are
        # therefore a straight projection of the rows in that sheet.  The
        # cover amount is retained only as reconciliation evidence; it must not
        # manufacture or allocate a target amount that is absent from details.
        quote_amount = max(detail_quote, Decimal("0")).quantize(
            KRW, rounding=ROUND_HALF_UP
        )
        negotiation = min(raw_negotiation, quote_amount)
        # A line whose historical target is above the received quote has no
        # additional negotiation room.  Keep that line at its received price,
        # then sum the line-level targets so one expensive target cannot offset
        # a saving identified on another line.
        target_amount = (quote_amount - negotiation).quantize(
            KRW, rounding=ROUND_HALF_UP
        )
        unallocated = Decimal("0")
        available = sum(target_by_raw[raw_id].status == "AVAILABLE" for raw_id in raw_ids)
        group = QuoteAnalysisEquipmentGroup(
            analysis_run_id=result.run_id,
            equipment_key=key,
            equipment_name=definition.name,
            source_kind=definition.source_kind,
            quote_amount=quote_amount,
            target_amount=target_amount,
            negotiation_amount=negotiation,
            unallocated_amount=unallocated,
            line_count=len(raw_ids),
            target_available_count=available,
            mapping_evidence_json=json.dumps(
                {
                    "sheet": definition.sheet,
                    "cover_quote_amount": (
                        None
                        if definition.cover_quote_amount is None
                        else str(definition.cover_quote_amount)
                    ),
                    "detail_quote_amount": str(detail_quote),
                    "detail_gap_amount": str(quote_amount - detail_quote),
                    "allocation_policy": "상세 시트 품목별 금액 직접 합산",
                },
                ensure_ascii=False,
                separators=(",", ":"),
            ),
        )
        session.add(group)
        session.flush()
        for raw_id in raw_ids:
            source_line = source_lines[raw_id]
            target_line = target_by_raw[raw_id]
            session.add(
                QuoteAnalysisEquipmentLine(
                    analysis_run_id=result.run_id,
                    equipment_group_id=group.id,
                    line_result_id=stored_by_raw[raw_id].id,
                    raw_item_id=raw_id,
                    quote_amount=max(source_line.quote_amount or Decimal("0"), Decimal("0")),
                    target_amount=target_line.target_amount,
                    negotiation_amount=_negotiation_amount(
                        source_line.quote_amount,
                        target_line.target_amount,
                    ),
                )
            )
        rows.append(group)
    session.flush()
    return tuple(rows)


def equipment_group_payloads(
    session: Session,
    analysis_run_id: int,
) -> list[dict[str, object]]:
    groups = list(
        session.scalars(
            select(QuoteAnalysisEquipmentGroup)
            .where(QuoteAnalysisEquipmentGroup.analysis_run_id == analysis_run_id)
            .order_by(QuoteAnalysisEquipmentGroup.id)
        )
    )
    lines = list(
        session.scalars(
            select(QuoteAnalysisEquipmentLine)
            .where(QuoteAnalysisEquipmentLine.analysis_run_id == analysis_run_id)
            .order_by(QuoteAnalysisEquipmentLine.raw_item_id)
        )
    )
    by_group: dict[int, list[dict[str, object]]] = defaultdict(list)
    for line in lines:
        by_group[line.equipment_group_id].append(
            {
                "raw_item_id": line.raw_item_id,
                "quote_amount": line.quote_amount,
                "target_amount": line.target_amount,
                "negotiation_amount": line.negotiation_amount,
            }
        )
    return [
        {
            "id": group.id,
            "key": group.equipment_key,
            "name": group.equipment_name,
            "source_kind": group.source_kind,
            "quote_amount": group.quote_amount,
            "target_amount": group.target_amount,
            "negotiation_amount": group.negotiation_amount,
            "unallocated_amount": group.unallocated_amount,
            "line_count": group.line_count,
            "target_available_count": group.target_available_count,
            "lines": by_group[group.id],
        }
        for group in groups
    ]


def _negotiation_amount(
    quote_amount: Decimal | None,
    target_amount: Decimal | None,
) -> Decimal:
    if quote_amount is None or target_amount is None:
        return Decimal("0")
    return max(quote_amount - min(quote_amount, target_amount), Decimal("0")).quantize(
        KRW, rounding=ROUND_HALF_UP
    )


def _equipment_definitions(result: AnalysisRunResult) -> tuple[EquipmentDefinition, ...]:
    paths = [line.source.path for line in result.analysis.lines if line.source.path]
    if not paths:
        return ()
    path = _resolve_source_path(Path(paths[0]))
    if not path.is_file() or path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        return ()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ()
    try:
        details: list[tuple[int, str, str]] = []
        for sheet_name in workbook.sheetnames:
            match = _DETAIL_SHEET.search(sheet_name) or _NUMBERED_SHEET.search(sheet_name)
            if match is None:
                continue
            sheet = workbook[sheet_name]
            name = _cell_text(sheet.cell(1, 3).value) or f"설비 {match.group(1)}"
            details.append((int(match.group(1)), sheet_name, name))
        if not details:
            return ()
        cover = next(
            (
                workbook[name]
                for name in workbook.sheetnames
                if "갑지" in name or "표지" in name
            ),
            workbook[workbook.sheetnames[0]],
        )
        scale = _cover_scale(cover)
        cover_amounts: dict[str, Decimal] = {}
        cover_names: dict[str, str] = {}
        for row_number in range(1, min(cover.max_row, 500) + 1):
            name = _cell_text(cover.cell(row_number, 3).value)
            amount = _decimal(cover.cell(row_number, 8).value)
            if name and amount is not None and amount > 0:
                scaled_amount = (amount * scale).quantize(KRW, rounding=ROUND_HALF_UP)
                normalized_name = _normalized_name(name)
                cover_names[normalized_name] = name
                cover_amounts[normalized_name] = scaled_amount

        def definition(detail: tuple[int, str, str]) -> EquipmentDefinition:
            number, sheet_name, detail_name = detail
            # The detail sheet is the calculation source.  When its title cell
            # matches a cover row, preserve the cover's official display name and
            # amount as reconciliation evidence only.
            normalized_detail_name = _normalized_name(detail_name)
            exact_amount = cover_amounts.get(normalized_detail_name)
            if exact_amount is not None:
                name = cover_names[normalized_detail_name]
                cover_amount = exact_amount
            else:
                name = detail_name
                cover_amount = None
            return EquipmentDefinition(
                key=f"equipment-{number}",
                name=name,
                sheet=sheet_name,
                cover_quote_amount=cover_amount,
                source_kind="COVER_SHEET",
            )

        return tuple(
            definition(detail)
            for detail in sorted(details)
        )
    finally:
        workbook.close()


def _resolve_source_path(path: Path) -> Path:
    if path.is_absolute():
        return path
    # Incoming submissions store a digest-relative path below the configured
    # submission folder. Historical corpus paths remain project-root relative.
    for root in (settings.submission_path, settings.project_root):
        root = root.resolve()
        candidate = (root / path).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            continue
        if candidate.is_file():
            return candidate
    return settings.project_root / path


def _cover_scale(sheet: object) -> Decimal:
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True):
        if any("천원" in str(value) for value in row if value is not None):
            return Decimal("1000")
    return Decimal("1")


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalized_name(value: str) -> str:
    return "".join(value.split()).casefold()


def _decimal(value: object) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        parsed = Decimal(str(value).replace(",", ""))
    except Exception:
        return None
    return parsed if parsed.is_finite() else None
