"""Recover auditable spreadsheet amount factors without evaluating formulas."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from app.core.config import settings


_CELL_REFERENCE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)", re.IGNORECASE)
_SAFE_PRODUCT = re.compile(
    r"^=\s*\$?[A-Z]{1,3}\$?\d+(?:\s*\*\s*\$?[A-Z]{1,3}\$?\d+)+\s*$",
    re.IGNORECASE,
)


def spreadsheet_amount_evidence(raw: object) -> dict[str, Any] | None:
    """Return same-row multiplication evidence when the workbook proves it.

    Only a plain multiplication of cells on the target row is accepted.  No
    arbitrary Excel formula is evaluated and cross-sheet formulas are left for
    review as composite calculations.
    """
    variant = getattr(raw, "__dict__", {}).get("source_variant")
    sheet_name = getattr(raw, "source_sheet", None)
    source_row = getattr(raw, "source_row", None)
    source_cells = getattr(raw, "source_cells", None)
    if variant is None or not sheet_name or not source_row or not source_cells:
        return None
    path = Path(variant.path)
    if not path.is_absolute():
        path = settings.quote_path / path
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        return None
    cell_matches = _CELL_REFERENCE.findall(source_cells)
    if not cell_matches:
        return None
    amount_coordinate = f"{cell_matches[-1][0]}{source_row}"
    formula_book = load_workbook(path, data_only=False, read_only=True)
    value_book = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in formula_book.sheetnames:
            return None
        formula_sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        formula = formula_sheet[amount_coordinate].value
        if not isinstance(formula, str) or not _SAFE_PRODUCT.fullmatch(formula):
            return None
        references = _CELL_REFERENCE.findall(formula)
        if len(references) < 2 or any(int(row) != source_row for _, row in references):
            return None
        factors: list[dict[str, str]] = []
        product = Decimal("1")
        for column, row_text in references:
            coordinate = f"{column.upper()}{row_text}"
            value = _decimal(value_sheet[coordinate].value)
            if value is None:
                return None
            product *= value
            factors.append(
                {
                    "label": _header_label(value_sheet, column.upper(), source_row),
                    "coordinate": coordinate,
                    "value": format(value, "f"),
                }
            )
        displayed = _decimal(value_sheet[amount_coordinate].value)
        if displayed is None:
            return None
        with localcontext() as context:
            context.prec = 64
            difference = product - displayed
            tolerance = max(Decimal("1"), abs(displayed) * Decimal("0.01"))
        return {
            "kind": "AMOUNT_CALCULATION",
            "formula": formula,
            "factors": factors,
            "calculated_amount": format(product, "f"),
            "displayed_amount": format(displayed, "f"),
            "difference_amount": format(difference, "f"),
            "difference_percent": (
                None
                if displayed == 0
                else format((difference / displayed * Decimal("100")).quantize(Decimal("0.0001")), "f")
            ),
            "tolerance_amount": format(tolerance, "f"),
            "matches": abs(difference) <= tolerance,
            "source": {"sheet": sheet_name, "row": source_row, "amount_cell": amount_coordinate},
        }
    finally:
        formula_book.close()
        value_book.close()


def _decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError):
        return None


def _header_label(sheet: object, column: str, source_row: int) -> str:
    for row_number in range(source_row - 1, max(0, source_row - 50), -1):
        value = sheet[f"{column}{row_number}"].value
        if isinstance(value, str) and value.strip():
            return value.strip()
    return column
